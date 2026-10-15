from __future__ import annotations

import argparse
import hashlib
import io
import json
import logging
import os
import re
import subprocess
import sys
import time
import traceback
from dataclasses import replace
from calendar import monthrange
from datetime import date, datetime, timedelta
from pathlib import Path

from PIL import Image
from PyQt5.QtCore import (
    QByteArray, QBuffer, QDate, QEvent, QEventLoop, QIODevice, QMimeData, QObject, QPoint,
    QRect, QRunnable, QThreadPool, QTimer, Qt, pyqtSignal,
)
from PyQt5.QtGui import QColor, QCursor, QFont, QIcon, QPainter, QPen, QPixmap
from PyQt5.QtNetwork import QLocalServer, QLocalSocket
from PyQt5.QtWidgets import (
    QApplication, QCalendarWidget, QCheckBox, QComboBox, QDateEdit, QDialog,
    QDialogButtonBox, QFileDialog, QFormLayout, QFrame, QGroupBox, QHBoxLayout,
    QLabel, QLineEdit, QListWidget, QListWidgetItem, QMainWindow,
    QMenu, QMessageBox, QPlainTextEdit, QPushButton, QStackedWidget,
    QScrollArea, QSpinBox, QSystemTrayIcon, QTableWidget, QTableWidgetItem, QTextBrowser,
    QVBoxLayout, QWidget,
)

from .core import Event, Store, parse_schedule_text, format_event_date
from .desktop import DEADLINE_ROLE, NoteWindow, DesktopListWindow, format_event_deadline
from .capture import CaptureNotepad, MimeTextEdit
from .review import ImportReviewDialog
from .layout import arrange_rectangles
from .urgency import event_urgency
from .importers import DocumentImportError, SUPPORTED_EXTENSIONS, extract_file_text
from .ocr import OcrEngine
from . import __version__

APP_NAME = "학교 알리미"
APP_VERSION = __version__
STARTUP_NAME = "SchoolAlimi"
TEST_NOTE_SOURCE = "알리미 포스트잇 동작 확인"
ALIGNMENTS = (("하단 정렬", "bottom"), ("상단 정렬", "top"),
              ("우측 정렬", "right"), ("좌측 정렬", "left"), ("자유 배치", "free"))


STYLE = """
QWidget { font-family: 'Malgun Gothic'; font-size: 10pt; color: #23382f; }
QMainWindow, QDialog { background: #f7f9f7; }
QPushButton { background: #ffffff; border: 1px solid #c9d6ce; border-radius: 6px;
  padding: 8px 14px; min-height: 18px; }
QPushButton:hover { background: #e8f0ea; }
QPushButton:disabled { color: #94a099; background: #edf0ed; }
QPushButton[primary="true"] { background: #285d47; color: white; border: 0; }
QLineEdit, QPlainTextEdit, QComboBox, QTextBrowser { background: white;
  border: 1px solid #ccd7cf; border-radius: 5px; padding: 7px; selection-background-color: #285d47; }
QCalendarWidget QWidget { background: white; }
QCalendarWidget QToolButton { color: #23382f; padding: 6px; }
QCalendarWidget QAbstractItemView { selection-background-color: #285d47; selection-color: white; }
QListWidget { background: white; border: 1px solid #d5dfd7; border-radius: 6px; padding: 5px; }
QListWidget::item { padding: 12px; }
QListWidget::item:selected { background: #285d47; color: white; border-radius: 4px; }
QTableWidget { background: white; border: 1px solid #d5dfd7; gridline-color: #e4eae5; }
QHeaderView::section { background: #eaf0eb; border: 0; padding: 7px; }
QLabel[muted="true"] { color: #62746a; }
QLabel[heading="true"] { font-size: 17pt; font-weight: 600; }
"""


def make_icon() -> QIcon:
    pixmap = QPixmap(64, 64)
    pixmap.fill(Qt.transparent)
    painter = QPainter(pixmap)
    painter.setRenderHint(QPainter.Antialiasing)
    painter.setPen(Qt.NoPen)
    painter.setBrush(QColor("#285d47"))
    painter.drawRoundedRect(2, 2, 60, 60, 13, 13)
    painter.setPen(QColor("white"))
    painter.setFont(QFont("Malgun Gothic", 29, QFont.Bold))
    painter.drawText(pixmap.rect(), Qt.AlignCenter, "알")
    painter.end()
    return QIcon(pixmap)


def default_data_dir() -> Path:
    return Path(os.environ.get("LOCALAPPDATA", str(Path.home()))) / "SchoolAlimi"


def frozen_install() -> None:
    """Install the new version before connecting to an existing app instance."""
    if not getattr(sys, "frozen", False):
        return
    if any(arg in sys.argv for arg in ("--data-dir", "--smoke-test", "--no-autostart")):
        return
    source = Path(sys.executable).resolve()
    try:
        from .updater import ensure_installed
        target = ensure_installed(source, default_data_dir(), APP_VERSION)
        if source != target.resolve():
            subprocess.Popen([str(target), *sys.argv[1:]], close_fds=True)
            raise SystemExit(0)
    except Exception as exc:
        bootstrap_app = QApplication.instance() or QApplication(sys.argv)
        QMessageBox.critical(None, APP_NAME, "새 버전으로 업데이트하지 못했습니다. 기존 일정은 보존됩니다.\n\n" + str(exc))
        raise SystemExit(1)


def set_startup(enabled: bool) -> None:
    import winreg
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER,
                         r"Software\Microsoft\Windows\CurrentVersion\Run") as key:
        if enabled:
            if getattr(sys, "frozen", False):
                command = f'"{sys.executable}" --background'
            else:
                pythonw = Path(sys.executable).with_name("pythonw.exe")
                command = f'"{pythonw}" "{Path(__file__).parents[1] / "main.py"}" --background'
            winreg.SetValueEx(key, STARTUP_NAME, 0, winreg.REG_SZ, command)
        else:
            try:
                winreg.DeleteValue(key, STARTUP_NAME)
            except FileNotFoundError:
                pass


def qimage_to_pil(image) -> Image.Image:
    data = QByteArray()
    buffer = QBuffer(data)
    buffer.open(QIODevice.WriteOnly)
    if not image.save(buffer, "PNG"):
        raise ValueError("이미지를 읽을 수 없습니다.")
    return Image.open(io.BytesIO(bytes(data))).convert("RGB")


class TaskSignals(QObject):
    result = pyqtSignal(object)
    error = pyqtSignal(str)
    finished = pyqtSignal()


class Task(QRunnable):
    def __init__(self, function):
        super().__init__()
        self.function = function
        self.signals = TaskSignals()

    def run(self):
        try:
            self.signals.result.emit(self.function())
        except Exception as exc:
            logging.error("Import failed: %s", type(exc).__name__)
            self.signals.error.emit(str(exc))
        finally:
            self.signals.finished.emit()


class CaptureOverlay(QWidget):
    captured = pyqtSignal(object)
    canceled = pyqtSignal()

    def __init__(self, screen):
        super().__init__(None, Qt.FramelessWindowHint | Qt.WindowStaysOnTopHint | Qt.Tool)
        self.screen = screen
        self.screenshot = screen.grabWindow(0)
        self.origin = None
        self.selection = QRect()
        self._finished = False
        self._close_prompt_active = False
        self.setGeometry(screen.geometry())
        self.setCursor(Qt.CrossCursor)
        self.setFocusPolicy(Qt.StrongFocus)
        self.setWindowTitle("일정 화면 캡처")
        self.close_button = QPushButton("닫기", self)
        self.close_button.setFixedSize(88, 34)
        self.close_button.setCursor(Qt.ArrowCursor)
        self.close_button.setStyleSheet(
            "QPushButton { background: white; color: #23382f; border: 1px solid #c9d6ce; "
            "border-radius: 6px; padding: 4px 12px; font: 10pt 'Malgun Gothic'; }"
            "QPushButton:hover { background: #e8f0ea; }"
        )
        self.close_button.clicked.connect(self.close)
        self._position_close_button()

    def _position_close_button(self):
        self.close_button.move(max(0, self.width() - self.close_button.width() - 18), 14)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, "close_button"):
            self._position_close_button()

    def _cancel_capture(self):
        if self._finished:
            return
        self._finished = True
        self.hide()
        self.canceled.emit()
        self.deleteLater()

    def closeEvent(self, event):
        if self._finished:
            event.accept()
            return
        application = QApplication.instance()
        if application and application.property("shutting_down"):
            self._cancel_capture()
            event.accept()
            return
        if self._close_prompt_active:
            event.ignore()
            return
        self._close_prompt_active = True
        try:
            answer = QMessageBox.warning(
                self, "화면 캡처 닫기", "화면 캡처를 닫을까요?\n선택한 영역은 가져오지 않습니다.",
                QMessageBox.Ok | QMessageBox.Cancel, QMessageBox.Cancel,
            )
        finally:
            self._close_prompt_active = False
        if answer == QMessageBox.Ok:
            self._cancel_capture()
            event.accept()
        else:
            event.ignore()

    def paintEvent(self, _):
        painter = QPainter(self)
        painter.drawPixmap(self.rect(), self.screenshot)
        painter.fillRect(self.rect(), QColor(0, 0, 0, 85))
        if not self.selection.isNull():
            painter.setClipRect(self.selection)
            painter.drawPixmap(self.rect(), self.screenshot)
            painter.setClipping(False)
            painter.setPen(QPen(QColor("#fff0ad"), 2))
            painter.drawRect(self.selection)
        painter.setPen(Qt.white)
        painter.setFont(QFont("Malgun Gothic", 12))
        painter.drawText(22, 35, "일정이 있는 영역을 드래그하세요 · Esc 취소")

    def mousePressEvent(self, event):
        if not self._finished and not self._close_prompt_active and event.button() == Qt.LeftButton:
            self.origin = event.pos()
            self.selection = QRect(self.origin, self.origin)

    def mouseMoveEvent(self, event):
        if not self._finished and not self._close_prompt_active and self.origin is not None:
            self.selection = QRect(self.origin, event.pos()).normalized().intersected(self.rect())
            self.update()

    def mouseReleaseEvent(self, event):
        if self._finished or self._close_prompt_active or event.button() != Qt.LeftButton or self.origin is None:
            return
        selection = self.selection.normalized()
        if selection.width() < 12 or selection.height() < 12:
            return
        scale_x = self.screenshot.width() / self.width()
        scale_y = self.screenshot.height() / self.height()
        physical = QRect(round(selection.x() * scale_x), round(selection.y() * scale_y),
                         round(selection.width() * scale_x), round(selection.height() * scale_y))
        image = qimage_to_pil(self.screenshot.copy(physical).toImage())
        self._finished = True
        self.hide()
        self.captured.emit(image)
        self.deleteLater()

    def keyPressEvent(self, event):
        if event.key() == Qt.Key_Escape:
            self._cancel_capture()


class EventDialog(QDialog):
    def __init__(self, event: Event | None = None, parent=None):
        super().__init__(parent)
        self.event = event or Event()
        self.setWindowTitle("일정 확인 / 수정" if event else "일정 직접 입력")
        self.resize(520, 560)
        layout = QVBoxLayout(self)
        form = QFormLayout()
        self.title_field = QLineEdit(self.event.title)
        initial_date = self.event.start_date[:7] if self.event.date_precision == "month" else self.event.start_date
        self.date_field = QLineEdit(initial_date or date.today().isoformat())
        self.date_field.setPlaceholderText("날짜: 2026-10-13 / 월: 2026-10")
        self.end_field = QLineEdit("" if self.event.date_precision == "month" else (self.event.end_date or ""))
        self.end_field.setPlaceholderText("기간 일정일 때만 입력")
        self.time_field = QLineEdit(self.event.time or "")
        self.time_field.setPlaceholderText("15:30 / 종일이면 비워두기")
        self.place_field = QLineEdit(self.event.place or "")
        self.notes_field = QPlainTextEdit(self.event.notes or "")
        self.notes_field.setMaximumHeight(80)
        form.addRow("일정 제목", self.title_field)
        form.addRow("일정 날짜 / 월", self.date_field)
        form.addRow("종료 날짜", self.end_field)
        form.addRow("시간", self.time_field)
        form.addRow("장소", self.place_field)
        form.addRow("메모", self.notes_field)
        layout.addLayout(form)
        source = QTextBrowser()
        source.setPlainText(self.event.source_text or "직접 입력한 일정")
        source.setMaximumHeight(140)
        layout.addWidget(QLabel("가져온 원문"))
        layout.addWidget(source)
        buttons = QDialogButtonBox(QDialogButtonBox.Save | QDialogButtonBox.Cancel)
        buttons.button(QDialogButtonBox.Save).setText("확정 저장")
        buttons.button(QDialogButtonBox.Cancel).setText("취소")
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def closeEvent(self, event):
        application = QApplication.instance()
        if application and application.property("shutting_down"):
            super().closeEvent(event)
            return
        answer = QMessageBox.warning(
            self, "일정 입력 창 닫기", "입력 창을 닫을까요?\n확정 저장하지 않은 수정 내용은 저장되지 않습니다.",
            QMessageBox.Ok | QMessageBox.Cancel, QMessageBox.Cancel,
        )
        if answer == QMessageBox.Ok:
            super().closeEvent(event)
        else:
            event.ignore()

    def accept(self):
        try:
            if not self.title_field.text().strip():
                raise ValueError("일정 제목을 입력해 주세요.")
            value = self.date_field.text().strip()
            precision = "month" if re.fullmatch(r"\d{4}-\d{2}", value) else "day"
            start = date.fromisoformat(value + "-01" if precision == "month" else value).isoformat()
            end = self.end_field.text().strip()
            if precision == "month":
                if end:
                    raise ValueError("월 일정은 종료 날짜를 비워주세요. 해당 월 전체에 표시됩니다.")
                monthly = date.fromisoformat(start)
                end = date(monthly.year, monthly.month, monthrange(monthly.year, monthly.month)[1]).isoformat()
            elif end:
                end = date.fromisoformat(end).isoformat()
                if end < start:
                    raise ValueError("종료일은 시작일 이후여야 합니다.")
            timing = self.time_field.text().strip()
            if timing:
                timing = datetime.strptime(timing, "%H:%M").strftime("%H:%M")
            self.event = replace(self.event, title=self.title_field.text().strip(),
                                 start_date=start, end_date=end, time=timing,
                                 place=self.place_field.text().strip(),
                                 notes=self.notes_field.toPlainText().strip(),
                                 status="done" if self.event.status == "done" else "open", date_precision=precision)
            super().accept()
        except ValueError as exc:
            QMessageBox.warning(self, "입력 확인", str(exc) or "날짜와 시간을 확인해 주세요.")


class MainWindow(QMainWindow):
    def __init__(self, store: Store, args, icon: QIcon):
        super().__init__()
        self.store, self.args, self.icon = store, args, icon
        self.notes: dict[int, NoteWindow] = {}
        self.desktop_list: DesktopListWindow | None = None
        self._dismissed_note_ids: set[int] = set()
        self._desktop_list_dismissed = False
        self._confirming_close = False
        self._daily_backup_pending = False
        self._daily_backup_retry_at = 0.0
        self._startup_refresh_remaining = 5 if args.background else 0
        self.desktop_mode = self.store.get_setting("desktop_mode", "sticky")
        if self.desktop_mode not in ("sticky", "list"):
            self.desktop_mode = "sticky"
        self._arranging_desktop = False
        self.jobs = set()
        self.pool = QThreadPool(self)
        self.pool.setMaxThreadCount(1)
        self.ocr = OcrEngine()
        self.overlay = None
        self.import_dialog = None
        self.import_queue = []
        self.capture_notepad = None
        self.capture_icon_position = None
        self.setWindowTitle(f"{APP_NAME} · {APP_VERSION}")
        self.setWindowIcon(icon)
        self.resize(1000, 730)
        self.setMinimumSize(760, 540)
        central = QWidget()
        self.setCentralWidget(central)
        outer = QVBoxLayout(central)
        outer.setContentsMargins(20, 20, 20, 15)
        heading = QLabel(APP_NAME)
        heading.setProperty("heading", True)
        heading_row = QHBoxLayout()
        heading_row.addWidget(heading, 1)
        self.close_button = QPushButton("닫기")
        self.close_button.setAccessibleName("관리 창 닫기")
        self.close_button.clicked.connect(self.close)
        heading_row.addWidget(self.close_button)
        outer.addLayout(heading_row)
        subtitle = QLabel("내 PC에만 저장 · 메모는 바탕화면에 독립적으로 표시")
        subtitle.setProperty("muted", True)
        outer.addWidget(subtitle)
        reminder_bar = QHBoxLayout()
        self.today_label = QLabel("오늘 날짜: " + date.today().isoformat())
        reminder_bar.addWidget(self.today_label, 1)
        reminder_bar.addWidget(QLabel("바탕화면 보기"))
        self.mode_combo = QComboBox()
        self.mode_combo.addItem("포스트잇", "sticky")
        self.mode_combo.addItem("목록", "list")
        self.mode_combo.setCurrentIndex(1 if self.desktop_mode == "list" else 0)
        self.mode_combo.currentIndexChanged.connect(self.change_desktop_mode)
        reminder_bar.addWidget(self.mode_combo)
        self.alignment_combo = self.make_alignment_combo()
        self.alignment_combo.setToolTip("바탕화면 알림의 정렬 방향을 선택합니다.")
        reminder_bar.addWidget(self.alignment_combo)
        self.align_bottom_button = QPushButton("정렬 적용")
        self.align_bottom_button.clicked.connect(self.apply_desktop_alignment)
        reminder_bar.addWidget(self.align_bottom_button)
        self.test_postit_button = QPushButton("포스트잇 동작 확인")
        self.test_postit_button.setToolTip("오늘의 테스트 일정 1개를 등록해 바탕화면 메모를 표시합니다.")
        self.test_postit_button.clicked.connect(self.show_test_postit)
        reminder_bar.addWidget(self.test_postit_button)
        outer.addLayout(reminder_bar)
        content = QHBoxLayout()
        outer.addLayout(content, 1)
        self.navigation = QListWidget()
        self.navigation.addItems(["달력", "일정 가져오기", "확인 필요", "완료 기록", "설정"])
        self.navigation.setFixedWidth(155)
        content.addWidget(self.navigation)
        self.pages = QStackedWidget()
        content.addWidget(self.pages, 1)
        self.make_calendar_page()
        self.make_import_page()
        self.review_table = self.make_event_table("review", "원문을 확인한 뒤 날짜와 내용을 확정해 주세요.")
        self.history_table = self.make_event_table("done", "완료한 일정은 보관됩니다. 실수로 완료한 일정은 되돌릴 수 있습니다.")
        self.make_settings_page()
        self.navigation.currentRowChanged.connect(self.pages.setCurrentIndex)
        self.navigation.setCurrentRow(0)
        self.status_label = QLabel("준비되었습니다.")
        self.status_label.setWordWrap(True)
        self.status_label.setProperty("muted", True)
        outer.addWidget(self.status_label)
        self.make_capture_notepad()
        self.text_input.textChanged.connect(lambda: self.draft_timer.start(300))
        self.make_tray()
        self.timer = QTimer(self)
        self.timer.timeout.connect(self.sync_notes)
        self.timer.start(10_000)
        self.startup_refresh_timer = QTimer(self)
        self.startup_refresh_timer.setInterval(2_000)
        self.startup_refresh_timer.timeout.connect(self.refresh_startup_reminders)
        self.midnight_timer = QTimer(self)
        self.midnight_timer.setSingleShot(True)
        self.midnight_timer.setTimerType(Qt.PreciseTimer)
        self.midnight_timer.timeout.connect(self.on_midnight)
        self.schedule_midnight()
        QApplication.instance().applicationStateChanged.connect(lambda _: self.sync_notes())
        QApplication.instance().aboutToQuit.connect(self.on_shutdown)
        QApplication.instance().commitDataRequest.connect(self.on_session_end)
        # A hidden background instance still needs a native window to receive
        # Windows' normal application-update shutdown notification.
        self.winId()
        self.refresh()
        self.apply_startup()
        self.sync_notes()
        if self._startup_refresh_remaining:
            self.startup_refresh_timer.start()

    def make_calendar_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        self.calendar = QCalendarWidget()
        self.calendar.setLocale(__import__("PyQt5.QtCore", fromlist=["QLocale"]).QLocale("ko_KR"))
        self.calendar.setGridVisible(False)
        self.calendar.setVerticalHeaderFormat(QCalendarWidget.NoVerticalHeader)
        self.calendar.selectionChanged.connect(self.refresh_agenda)
        layout.addWidget(self.calendar, 3)
        toolbar = QHBoxLayout()
        self.day_label = QLabel()
        toolbar.addWidget(self.day_label, 1)
        today_button = QPushButton("오늘로 이동")
        today_button.clicked.connect(lambda: self.calendar.setSelectedDate(QDate.currentDate()))
        toolbar.addWidget(today_button)
        add = QPushButton("직접 일정 추가")
        add.setProperty("primary", True)
        add.clicked.connect(self.add_manual)
        toolbar.addWidget(add)
        layout.addLayout(toolbar)
        self.agenda = QListWidget()
        self.agenda.itemDoubleClicked.connect(lambda item: self.edit_event(item.data(Qt.UserRole)))
        self.agenda.currentItemChanged.connect(self.update_note_reason)
        layout.addWidget(self.agenda, 2)
        self.note_reason = QLabel()
        self.note_reason.setTextFormat(Qt.PlainText)
        self.note_reason.setWordWrap(True)
        self.note_reason.setProperty("muted", True)
        layout.addWidget(self.note_reason)
        self.edit_selected_button = QPushButton("선택 일정 확인 / 수정")
        self.edit_selected_button.setEnabled(False)
        self.edit_selected_button.clicked.connect(self.edit_selected_event)
        layout.addWidget(self.edit_selected_button)
        self.pages.addWidget(page)

    def make_import_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        row = QHBoxLayout()
        paste = QPushButton("클립보드 붙여넣기")
        paste.clicked.connect(self.paste_clipboard)
        self.capture_button = QPushButton("화면 영역 캡처")
        self.capture_button.clicked.connect(self.capture_screen)
        self.file_button = QPushButton("파일에서 일정 가져오기")
        self.file_button.clicked.connect(self.import_files)
        sample = QPushButton("오늘 공지 예시")
        sample.clicked.connect(self.fill_sample)
        for button in (paste, self.capture_button, self.file_button, sample):
            row.addWidget(button)
        layout.addLayout(row)
        self.text_input = MimeTextEdit()
        self.text_input.setPlainText(self.store.get_setting("recognized_text_draft", ""))
        self.import_preview = None
        self.text_input.payload_received.connect(lambda payload: self.receive_import_payload(payload, "manager"))
        self.text_input.setPlaceholderText("공지를 붙여넣거나 위에서 파일을 선택해 주세요.\n파일에서 읽은 원문도 이곳에 표시됩니다.\n예: 2026년 10월 13일(화) 15:30 교직원 협의회, 장소: 2층 회의실")
        layout.addWidget(self.text_input, 2)
        self.analyze_button = QPushButton("일정 찾고 확인")
        self.analyze_button.setProperty("primary", True)
        self.analyze_button.clicked.connect(self.import_text)
        layout.addWidget(self.analyze_button)
        support = QLabel("붙여넣거나 파일을 끌어 놓으면 확인 창을 엽니다. 내용을 수정한 뒤 ‘등록’을 눌러 저장하세요.\n지원: TXT · PDF · HWPX · HWP · DOCX · XLSX · PNG/JPG\n달만 적힌 일정은 ‘9월’처럼 월로 표시하고, 해당 월이 시작되면 알림을 띄웁니다.")
        support.setWordWrap(True)
        support.setProperty("muted", True)
        layout.addWidget(support)
        self.import_result = QTextBrowser()
        self.import_result.setPlainText("가져온 원문과 처리 결과가 이곳에 표시됩니다.")
        layout.addWidget(self.import_result, 1)
        self.pages.addWidget(page)

    def make_event_table(self, mode, hint):
        page = QWidget()
        layout = QVBoxLayout(page)
        label = QLabel(hint)
        label.setWordWrap(True)
        layout.addWidget(label)
        table = QTableWidget(0, 4)
        table.setHorizontalHeaderLabels(["제목", "날짜", "시간", "출처"])
        table.setSelectionBehavior(QTableWidget.SelectRows)
        table.setEditTriggers(QTableWidget.NoEditTriggers)
        table.horizontalHeader().setStretchLastSection(True)
        table.cellDoubleClicked.connect(lambda row, _: self.edit_event(table.item(row, 0).data(Qt.UserRole)))
        layout.addWidget(table)
        button = QPushButton("선택 일정 확인 / 확정" if mode == "review" else "선택 일정 완료 취소")
        button.clicked.connect(lambda: self.table_action(table, mode))
        layout.addWidget(button)
        self.pages.addWidget(page)
        return table

    def make_settings_page(self):
        page = QWidget()
        layout = QVBoxLayout(page)
        self.autostart = QCheckBox("Windows 로그인 시 자동으로 실행")
        self.autostart.setChecked(str(self.store.get_setting("autostart", "True")).lower() in ("true", "1", "yes"))
        self.autostart.toggled.connect(self.change_startup)
        layout.addWidget(self.autostart)
        window_group = QGroupBox("바탕화면에 표시할 미완료 일정")
        window_layout = QVBoxLayout(window_group)
        self.reminder_unlimited = QCheckBox("날짜 제한 없이 모든 미완료 일정 표시")
        window_layout.addWidget(self.reminder_unlimited)
        window_form = QFormLayout()
        self.reminder_past_days = QSpinBox()
        self.reminder_past_days.setRange(0, 36500)
        self.reminder_past_days.setSuffix("일 전까지")
        self.reminder_future_days = QSpinBox()
        self.reminder_future_days.setRange(0, 36500)
        self.reminder_future_days.setSuffix("일 후까지")
        window_form.addRow("오늘 기준 지난 일정", self.reminder_past_days)
        window_form.addRow("오늘 기준 앞으로의 일정", self.reminder_future_days)
        window_layout.addLayout(window_form)
        self.reminder_window_hint = QLabel("마감 날짜를 오늘과 비교합니다. 월만 있는 일정은 해당 월과 표시 기간이 겹치면 표시합니다.\n표시 기간 밖이어도 삭제하거나 완료 처리하지 않습니다. 달력에서 계속 확인할 수 있습니다.")
        self.reminder_window_hint.setWordWrap(True)
        self.reminder_window_hint.setProperty("muted", True)
        window_layout.addWidget(self.reminder_window_hint)
        self.sync_reminder_controls()
        self.reminder_unlimited.toggled.connect(self.change_reminder_window)
        self.reminder_past_days.valueChanged.connect(self.change_reminder_window)
        self.reminder_future_days.valueChanged.connect(self.change_reminder_window)
        layout.addWidget(window_group)
        row = QHBoxLayout()
        row.addWidget(QLabel("알림 글자 크기"))
        self.font_combo = QComboBox()
        self.font_combo.addItems(["11", "12", "14", "16", "18"])
        self.font_combo.setCurrentText(str(self.store.get_setting("font_size", 12)))
        self.font_combo.currentTextChanged.connect(self.change_font)
        row.addWidget(self.font_combo)
        row.addStretch()
        layout.addLayout(row)
        alignment_row = QHBoxLayout()
        alignment_row.addWidget(QLabel("바탕화면 알림 정렬"))
        self.settings_alignment_combo = self.make_alignment_combo()
        alignment_row.addWidget(self.settings_alignment_combo)
        alignment_row.addStretch()
        layout.addLayout(alignment_row)
        desktop_hint = QLabel("접힌 바탕화면 알림은 다른 창 아래에 유지됩니다. 포스트잇 제목을 눌러 세부 내용을 펼치면 다른 창 앞에 표시됩니다.\n상단·하단·좌측·우측 정렬을 선택하면 바로 적용하고 저장합니다. 이동 영역을 드래그하면 자유 배치로 바뀝니다.")
        desktop_hint.setWordWrap(True)
        desktop_hint.setProperty("muted", True)
        layout.addWidget(desktop_hint)
        self.ocr_info = QLabel(self.ocr.diagnostics())
        self.ocr_info.setWordWrap(True)
        layout.addWidget(self.ocr_info)
        path_info = QLabel("데이터 저장 폴더:\n" + str(self.args.data_dir))
        path_info.setWordWrap(True)
        layout.addWidget(path_info)
        open_folder = QPushButton("데이터 폴더 열기")
        open_folder.clicked.connect(lambda: os.startfile(str(self.args.data_dir)))
        layout.addWidget(open_folder)
        backup = QPushButton("일정 데이터 백업")
        backup.clicked.connect(self.backup_data)
        layout.addWidget(backup)
        restore = QPushButton("백업 파일에서 복원")
        restore.clicked.connect(self.restore_data)
        layout.addWidget(restore)
        reset_group = QGroupBox("일정 초기화")
        reset_layout = QVBoxLayout(reset_group)
        reset_hint = QLabel(
            "완료·확인 필요 일정을 포함해 삭제합니다. 기간 초기화는 선택 기간과 겹치는 일정과 월 일정을 대상으로 합니다.\n"
            "날짜 없는 일정은 전체 초기화에만 포함됩니다. 초기화 전에 자동 백업을 만듭니다."
        )
        reset_hint.setWordWrap(True)
        reset_layout.addWidget(reset_hint)
        reset_dates = QFormLayout()
        today = QDate.currentDate()
        self.reset_from_date = QDateEdit(QDate(today.year(), today.month(), 1))
        self.reset_to_date = QDateEdit(QDate(today.year(), today.month(), today.daysInMonth()))
        for control in (self.reset_from_date, self.reset_to_date):
            control.setDisplayFormat("yyyy-MM-dd")
            control.setCalendarPopup(True)
            control.setDateRange(QDate(1, 1, 1), QDate(9999, 12, 31))
        reset_dates.addRow("시작일", self.reset_from_date)
        reset_dates.addRow("종료일", self.reset_to_date)
        reset_layout.addLayout(reset_dates)
        reset_buttons = QHBoxLayout()
        self.reset_range_button = QPushButton("지정 기간 일정 초기화")
        self.reset_range_button.clicked.connect(self.reset_schedule_range)
        reset_buttons.addWidget(self.reset_range_button)
        self.reset_all_button = QPushButton("전체 일정 초기화")
        self.reset_all_button.setStyleSheet("color: #a12b21;")
        self.reset_all_button.clicked.connect(self.reset_all_schedules)
        reset_buttons.addWidget(self.reset_all_button)
        reset_layout.addLayout(reset_buttons)
        self.reset_result_label = QLabel("초기화한 일정은 자동 백업에서 복원할 수 있습니다. 설정과 메모 초안은 유지합니다.")
        self.reset_result_label.setTextFormat(Qt.PlainText)
        self.reset_result_label.setWordWrap(True)
        reset_layout.addWidget(self.reset_result_label)
        layout.addWidget(reset_group)
        privacy = QLabel("인터넷 없이 작동합니다. 원문, 문자 인식 결과, 일정과 완료 기록을 외부로 전송하지 않습니다.\n로컬 파일은 Windows 계정의 접근 권한으로 보호됩니다.")
        privacy.setWordWrap(True)
        privacy.setProperty("muted", True)
        layout.addWidget(privacy)
        layout.addStretch()
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.NoFrame)
        scroll.setWidget(page)
        self.pages.addWidget(scroll)

    def sync_reminder_controls(self):
        unlimited, past_days, future_days = self.store.get_reminder_window()
        controls = (self.reminder_unlimited, self.reminder_past_days, self.reminder_future_days)
        for control in controls:
            control.blockSignals(True)
        self.reminder_unlimited.setChecked(unlimited)
        self.reminder_past_days.setValue(past_days)
        self.reminder_future_days.setValue(future_days)
        self.reminder_past_days.setEnabled(not unlimited)
        self.reminder_future_days.setEnabled(not unlimited)
        for control in controls:
            control.blockSignals(False)

    def change_reminder_window(self, _value=None):
        unlimited = self.reminder_unlimited.isChecked()
        self.store.set_setting("reminder_unlimited", str(unlimited).lower())
        self.store.set_setting("reminder_past_days", self.reminder_past_days.value())
        self.store.set_setting("reminder_future_days", self.reminder_future_days.value())
        self.reminder_past_days.setEnabled(not unlimited)
        self.reminder_future_days.setEnabled(not unlimited)
        self.refresh()
        self.status_label.setText("바탕화면 표시 기간을 적용했습니다. " + self.reminder_window_description())

    def reminder_window_description(self):
        unlimited, past_days, future_days = self.store.get_reminder_window()
        if unlimited:
            return "날짜 제한 없이 모든 미완료 일정을 표시합니다."
        return f"오늘 기준 {past_days}일 전부터 {future_days}일 후까지 표시합니다."

    def make_capture_notepad(self):
        self.capture_notepad = CaptureNotepad(
            self.store.get_setting("capture_draft", ""),
            int(self.store.get_setting("font_size", 12)))
        self.capture_notepad.setWindowIcon(self.icon)
        self.draft_timer = QTimer(self)
        self.draft_timer.setSingleShot(True)
        self.draft_timer.timeout.connect(self.save_capture_draft)
        self.capture_notepad.draft_changed.connect(lambda _text: self.draft_timer.start(300))
        self.capture_notepad.payload_received.connect(self.receive_import_payload)
        self.capture_notepad.review_requested.connect(self.review_capture_draft)
        self.capture_notepad.moved.connect(self.save_capture_position)
        self.capture_notepad.view_changed.connect(self.capture_view_changed)
        self.capture_notepad.dismissed_by_user.connect(self.save_capture_draft)
        x = self.store.get_setting("capture_x", "")
        y = self.store.get_setting("capture_y", "")
        wanted = QPoint(int(x), int(y)) if str(x).lstrip("-").isdigit() and str(y).lstrip("-").isdigit() else None
        screen = next((s for s in QApplication.screens() if wanted and s.availableGeometry().contains(wanted)), QApplication.primaryScreen())
        area = screen.availableGeometry()
        left = wanted.x() if wanted else area.right() - self.capture_notepad.width() - 11
        top = wanted.y() if wanted else area.top() + 20
        self.capture_notepad.move(max(area.left(), min(left, area.right() - self.capture_notepad.width() + 1)),
                                  max(area.top(), min(top, area.bottom() - self.capture_notepad.height() + 1)))
        self.capture_icon_position = QPoint(self.capture_notepad.pos())
        self.show_capture_notepad()

    def show_capture_notepad(self, expand=False):
        if self.capture_notepad:
            if expand:
                self.capture_notepad.dismissed = False
                self.capture_notepad.expand_notepad()
            if self.capture_notepad.dismissed:
                return
            self.capture_notepad.show()
            self.capture_notepad.keep_below_windows()

    def restore_desktop_tools(self):
        self._dismissed_note_ids.clear()
        self._desktop_list_dismissed = False
        for note in self.notes.values():
            note.dismissed = False
            note.show()
        if self.desktop_list:
            self.desktop_list.dismissed = False
        self.sync_notes()
        self.refresh_agenda()
        self.status_label.setText("닫았던 바탕화면 알림을 다시 표시했습니다.")

    def dismiss_note(self, event_id):
        self._dismissed_note_ids.add(event_id)
        self.refresh_agenda()
        # closeEvent hides the window after this signal is delivered.
        QTimer.singleShot(0, self.arrange_desktop_bottom)

    def dismiss_desktop_list(self):
        self._desktop_list_dismissed = True
        self.refresh_agenda()

    def capture_view_changed(self, collapsed):
        note = self.capture_notepad
        if collapsed:
            if self.capture_icon_position is not None:
                note.move(self.capture_icon_position)
        else:
            self.capture_icon_position = QPoint(note.pos())
        anchor = self.capture_icon_position or note.pos()
        screen = next((screen for screen in QApplication.screens() if screen.availableGeometry().contains(anchor)), QApplication.primaryScreen())
        area = screen.availableGeometry()
        note.move(max(area.left(), min(note.x(), area.right() - note.width() + 1)),
                  max(area.top(), min(note.y(), area.bottom() - note.height() + 1)))
        if collapsed:
            self.save_capture_position(note.x(), note.y())
        note.keep_below_windows()

    def save_capture_draft(self):
        if self.capture_notepad:
            self.store.set_setting("capture_draft", self.capture_notepad.editor.toPlainText())
            self.store.set_setting("recognized_text_draft", self.text_input.toPlainText())

    def save_capture_position(self, x, y):
        self.capture_icon_position = QPoint(x, y)
        self.store.set_setting("capture_x", x)
        self.store.set_setting("capture_y", y)

    def review_capture_draft(self):
        text = self.capture_notepad.editor.toPlainText().strip()
        if not text:
            self.capture_notepad.status_label.setText("일정 내용을 입력하거나 붙여넣어 주세요.")
            return
        self.save_capture_draft()
        self.receive_import_text(text, "일정 메모장")

    def make_tray(self):
        self.tray = QSystemTrayIcon(self.icon, self)
        self.tray.setToolTip(APP_NAME)
        menu = QMenu()
        menu.addAction("달력 열기", self.show_manager)
        menu.addAction("일정 메모장 열기", lambda _checked=False: self.show_capture_notepad(expand=True))
        menu.addAction("바탕화면 알림 다시 열기", self.restore_desktop_tools)
        menu.addAction("일정 직접 추가", self.add_manual)
        menu.addAction("화면 영역 캡처", self.capture_screen)
        menu.addSeparator()
        menu.addAction("알리미 종료", self.request_quit)
        self.tray.setContextMenu(menu)
        self.tray.activated.connect(lambda reason: self.show_manager() if reason in (QSystemTrayIcon.Trigger, QSystemTrayIcon.DoubleClick) else None)
        if not self.args.smoke_test:
            self.tray.show()

    def apply_startup(self):
        if self.args.no_autostart or self.args.smoke_test:
            return
        try:
            set_startup(self.autostart.isChecked())
        except OSError:
            self.status_label.setText("학교 PC 정책으로 자동 시작을 등록하지 못했습니다. 설정에서 다시 시도할 수 있습니다.")

    def change_startup(self, enabled):
        try:
            if not self.args.no_autostart and not self.args.smoke_test:
                set_startup(enabled)
            self.store.set_setting("autostart", enabled)
            self.status_label.setText("자동 시작 설정을 저장했습니다.")
        except Exception as exc:
            self.show_error(str(exc))

    def change_font(self, value):
        self.store.set_setting("font_size", int(value))
        self.capture_notepad.editor.setFont(QFont("Malgun Gothic", int(value)))
        for note in self.notes.values():
            note.set_font_size(int(value))
        if self.desktop_list:
            self.desktop_list.list_widget.setFont(QFont("Malgun Gothic", int(value)))
            self.desktop_list.details_label.setFont(QFont("Malgun Gothic", max(10, int(value) - 2)))
            self.desktop_list.adjustSize()
        self.arrange_desktop_bottom()

    def change_desktop_mode(self, _index):
        self.desktop_mode = self.mode_combo.currentData()
        self.store.set_setting("desktop_mode", self.desktop_mode)
        self.sync_notes()
        self.refresh_agenda()

    def make_alignment_combo(self):
        combo = QComboBox()
        for label, value in ALIGNMENTS:
            combo.addItem(label, value)
        value = self.store.get_setting("desktop_layout", "bottom")
        index = combo.findData(value)
        combo.setCurrentIndex(max(0, index))
        combo.currentIndexChanged.connect(lambda _index: self.set_desktop_alignment(combo.currentData()))
        return combo

    def sync_alignment_controls(self, alignment):
        for name in ("alignment_combo", "settings_alignment_combo"):
            combo = getattr(self, name, None)
            if combo:
                combo.blockSignals(True)
                combo.setCurrentIndex(max(0, combo.findData(alignment)))
                combo.blockSignals(False)

    def set_desktop_alignment(self, alignment):
        if alignment not in {value for _label, value in ALIGNMENTS}:
            alignment = "bottom"
        self.store.set_setting("desktop_layout", alignment)
        if alignment != "free":
            self.store.set_setting("desktop_alignment", alignment)
        self.sync_alignment_controls(alignment)
        self.arrange_desktop()
        if hasattr(self, "status_label"):
            label = next(label for label, value in ALIGNMENTS if value == alignment)
            self.status_label.setText("바탕화면 알림: " + label + "을 적용했습니다.")

    def change_desktop_alignment(self, _index):
        self.set_desktop_alignment(self.alignment_combo.currentData())

    def apply_desktop_alignment(self):
        alignment = self.alignment_combo.currentData()
        if alignment == "free":
            alignment = self.store.get_setting("desktop_alignment", "bottom")
        self.set_desktop_alignment(alignment)

    def align_desktop_bottom(self):
        self.set_desktop_alignment("bottom")

    def desktop_size_changed(self):
        if not self._arranging_desktop:
            self.arrange_desktop_bottom()

    def save_list_position(self, x, y):
        self.store.set_setting("desktop_layout", "free")
        self.sync_alignment_controls("free")
        self.store.set_setting("list_x", x)
        self.store.set_setting("list_y", y)

    def show_manager(self):
        self.showNormal()
        self.raise_()
        self.activateWindow()

    def closeEvent(self, event):
        if QApplication.instance().property("shutting_down"):
            event.accept()
            return
        event.ignore()
        if self._confirming_close:
            return
        self._confirming_close = True
        try:
            answer = QMessageBox.warning(
                self, "관리 창 닫기",
                "관리 창을 닫을까요?\n알리미는 트레이에서 계속 실행됩니다. 트레이의 '달력 열기'에서 다시 열 수 있습니다.",
                QMessageBox.Ok | QMessageBox.Cancel, QMessageBox.Cancel,
            )
        finally:
            self._confirming_close = False
        if answer == QMessageBox.Ok:
            self.hide()

    def request_quit(self):
        if self.jobs:
            QMessageBox.information(self, "자료 분석 중", "자료 분석이 끝난 뒤 종료해 주세요.")
            return
        try:
            # Schedules are committed when registered or edited. Flush the
            # draft before exiting, including text still in its debounce timer.
            self.save_capture_draft()
            icon_position = self.capture_icon_position or self.capture_notepad.pos()
            self.save_capture_position(icon_position.x(), icon_position.y())
        except Exception as exc:
            self.show_error("종료 전 메모장 상태를 로컬에 저장하지 못해 종료하지 않았습니다.\n" + str(exc))
            return
        QApplication.instance().setProperty("shutting_down", True)
        QApplication.instance().quit()

    def refresh(self):
        self.refresh_agenda()
        events = self.store.list_events()
        for table, status in ((self.review_table, "review"), (self.history_table, "done")):
            current = table.item(table.currentRow(), 0) if table.currentRow() >= 0 else None
            selected_id = current.data(Qt.UserRole) if current else None
            table.clearSelection()
            table.setCurrentCell(-1, -1)
            selected = [event for event in events if event.status == status]
            table.setRowCount(len(selected))
            for row, event in enumerate(selected):
                for col, value in enumerate((event.title, format_event_date(event, include_year=True), event.time or "", event.source or "")):
                    item = QTableWidgetItem(value)
                    item.setData(Qt.UserRole, event.id)
                    table.setItem(row, col, item)
                if event.id == selected_id:
                    table.setCurrentCell(row, 0)
            table.resizeColumnsToContents()
        self.mark_calendar(events)
        self.sync_notes()

    def mark_calendar(self, events):
        from PyQt5.QtGui import QTextCharFormat
        self.calendar.setDateTextFormat(QDate(), QTextCharFormat())
        for event in events:
            if not event.start_date:
                continue
            initial = date.fromisoformat(event.start_date)
            days = range(1, monthrange(initial.year, initial.month)[1] + 1) if event.date_precision == "month" else [initial.day]
            for day_number in days:
                day = QDate(initial.year, initial.month, day_number)
                format_ = QTextCharFormat()
                format_.setFontWeight(QFont.Bold)
                format_.setForeground(QColor("#936d15" if event.status == "review" else "#285d47"))
                self.calendar.setDateTextFormat(day, format_)

    def refresh_agenda(self):
        if not hasattr(self, "agenda"):
            return
        selected = self.calendar.selectedDate().toString("yyyy-MM-dd")
        self.day_label.setText(selected + " 일정")
        current_item = self.agenda.currentItem()
        previous_id = current_item.data(Qt.UserRole) if current_item else None
        self.agenda.clear()
        events = [event for event in self.store.list_events() if event.start_date == selected or
                  (event.start_date and event.end_date and event.start_date <= selected <= event.end_date)]
        eligible_ids = {event.id for event in self.store.reminder_events()}
        for event in events:
            tool_closed = (self._desktop_list_dismissed if self.desktop_mode == "list" else
                           event.id in self._dismissed_note_ids)
            state = ("확인 필요 · 알림 보류" if event.status == "review" else
                     "완료" if event.status == "done" else
                     "미완료 · 설정한 표시 기간 밖" if event.id not in eligible_ids else
                     "미완료 · 바탕화면 도구 닫힘" if tool_closed else
                     "미완료 · 바탕화면 알림 표시 중")
            timing = format_event_date(event) if event.date_precision == "month" else (event.time or '종일')
            item = QListWidgetItem(f"{timing}  {event.title}\n{event.place or event.source}  ·  {state}")
            item.setData(Qt.UserRole, event.id)
            self.agenda.addItem(item)
            if event.id == previous_id:
                self.agenda.setCurrentItem(item)
        if not events:
            item = QListWidgetItem("등록된 일정이 없습니다.")
            item.setFlags(Qt.NoItemFlags)
            self.agenda.addItem(item)
        elif self.agenda.currentRow() < 0:
            self.agenda.setCurrentRow(0)
        self.update_note_reason(self.agenda.currentItem())

    def edit_selected_event(self):
        item = self.agenda.currentItem()
        if item and item.data(Qt.UserRole) is not None:
            self.edit_event(item.data(Qt.UserRole))

    def update_note_reason(self, item=None, _previous=None):
        if not hasattr(self, "note_reason"):
            return
        id_ = item.data(Qt.UserRole) if item else None
        event = self.store.get_event(id_) if id_ is not None else None
        self.edit_selected_button.setEnabled(event is not None)
        if event is None:
            text = self.reminder_window_description() + " 마감 3일 전부터 색이 진해지고 당일은 주황색이 됩니다. 월만 있는 일정은 기본색을 유지합니다."
        elif event.status == "review":
            reason = event.notes or "원문의 날짜와 내용을 확인해 주세요."
            text = "포스트잇 표시 보류: 확인 필요 상태입니다. " + reason
            text += "\n아래 ‘선택 일정 확인 / 수정’에서 확정 저장해 주세요."
        elif event.status == "done":
            text = "완료한 일정이라 포스트잇을 표시하지 않습니다. 완료 기록에서 완료 취소할 수 있습니다."
        elif event.id not in {item.id for item in self.store.reminder_events()}:
            text = f"마감: {format_event_deadline(event, include_year=True)}. 일정은 저장되어 있으며, 설정한 바탕화면 표시 기간 밖이라 숨겨져 있습니다. " + self.reminder_window_description() + " 설정에서 기간을 늘리거나 ‘날짜 제한 없이’ 표시를 선택하세요."
        elif ((self.desktop_mode == "list" and self._desktop_list_dismissed) or
              (self.desktop_mode == "sticky" and event.id in self._dismissed_note_ids)):
            text = "바탕화면 도구를 닫아 이 일정이 숨겨져 있습니다. 일정은 미완료 상태로 저장되어 있습니다. 트레이의 '바탕화면 알림 다시 열기'에서 다시 표시할 수 있습니다."
        else:
            text = "이 일정은 바탕화면에 표시됩니다. " + event_urgency(event).label + ". 관리 창을 닫아도 알림은 남습니다. 접힌 알림은 다른 창 아래에, 펼친 포스트잇은 다른 창 앞에 표시됩니다."
        self.note_reason.setText(text)

    def show_test_postit(self):
        try:
            today = date.today().isoformat()
            event = next((e for e in self.store.list_events() if e.source == TEST_NOTE_SOURCE and e.start_date == today), None)
            if event is None:
                event = self.store.upsert_event(Event(
                    title="포스트잇 동작 확인 (테스트)", start_date=today,
                    place="바탕화면", notes="메모가 보이면 완료를 눌러주세요.",
                    source=TEST_NOTE_SOURCE, status="open"))
                if event.source != TEST_NOTE_SOURCE:
                    raise ValueError("같은 이름의 일정이 이미 있습니다. 직접 일정 추가로 오늘 일정을 등록해 주세요.")
            elif event.status == "done":
                self.store.reopen_event(event.id)
            elif event.status == "review":
                self.store.upsert_event(replace(event, status="open"))
            self._dismissed_note_ids.discard(int(event.id))
            existing_note = self.notes.get(int(event.id))
            if existing_note:
                existing_note.dismissed = False
            if self.desktop_mode == "list":
                self._desktop_list_dismissed = False
                if self.desktop_list:
                    self.desktop_list.dismissed = False
            self.calendar.setSelectedDate(QDate.currentDate())
            self.navigation.setCurrentRow(0)
            self.refresh()
            note = self.notes.get(int(event.id))
            if note:
                note.show()
                note.keep_below_windows()
                self.store.update_position(event.id, note.x(), note.y())
            self.status_label.setText("테스트 알림을 바탕화면에 표시했습니다. 제목을 눌러 펼친 뒤 완료를 눌러주세요.")
        except Exception as exc:
            self.show_error(str(exc))

    def safe_position(self, note, event, index):
        screens = QApplication.screens()
        wanted = QPoint(event.x, event.y) if event.x is not None and event.y is not None else None
        screen = next((s for s in screens if wanted is not None and s.availableGeometry().contains(wanted)), QApplication.primaryScreen())
        area = screen.availableGeometry()
        if wanted is None:
            wanted = QPoint(area.left() + 12, area.bottom() - note.height() - 11)
            # Newly eligible reminders need their own position in free mode.
            # Keep existing manually placed cards where the user left them.
            cell_height = max(105, note.height())
            columns = max(1, (area.width() - 24 + 10) // (note.width() + 10))
            rows = max(1, (area.height() - 24 + 10) // (cell_height + 10))
            alignment = self.store.get_setting("desktop_alignment", "bottom")
            positions = arrange_rectangles(area, [(note.width(), cell_height)] * (columns * rows), alignment)
            occupied = [other.geometry().adjusted(-5, -5, 5, 5) for other in self.notes.values()
                        if other is not note and other.isVisible()]
            for position in positions:
                rectangle = QRect(position, note.size())
                if area.contains(rectangle) and not any(rectangle.intersects(other) for other in occupied):
                    wanted = position
                    break
        note.move(max(area.left(), min(wanted.x(), area.right() - note.width() + 1)),
                  max(area.top(), min(wanted.y(), area.bottom() - note.height() + 1)))
        if event.x is None and event.y is None and event.id is not None and self.notes.get(int(event.id)) is note:
            self.store.update_position(event.id, note.x(), note.y())

    def arrange_desktop_bottom(self):
        self.arrange_desktop()

    def arrange_desktop(self):
        if self._arranging_desktop:
            return
        self._arranging_desktop = True
        try:
            alignment = self.store.get_setting("desktop_layout", "bottom")
            if alignment not in {value for _label, value in ALIGNMENTS}:
                alignment = "bottom"
            automatic = alignment != "free"
            area = QApplication.primaryScreen().availableGeometry()
            if self.desktop_mode == "list":
                window = self.desktop_list
                if window is None or window.drag_offset is not None:
                    return
                window.resize(window.width(), min(window.height(), max(300, area.height() - 24)))
                x = area.right() - window.width() - 11 if alignment != "left" else area.left() + 12
                y = area.bottom() - window.height() - 11 if alignment == "bottom" else area.top() + 12
                if not automatic:
                    x = int(self.store.get_setting("list_x", x))
                    y = int(self.store.get_setting("list_y", y))
                self.safe_position(window, Event(x=x, y=y), 0)
                window.keep_below_windows()
                return
            ordered = [self.notes[int(e.id)] for e in self.store.reminder_events()
                       if int(e.id) in self.notes and int(e.id) not in self._dismissed_note_ids]
            if not ordered:
                return
            if not automatic:
                for index, note in enumerate(ordered):
                    if note.drag_offset is None:
                        self.safe_position(note, self.store.get_event(note.event.id) or note.event, index)
                        note.keep_below_windows()
                return
            for note in ordered:
                if note.expanded:
                    note.details_scroll.setFixedHeight(220)
                    note.adjustSize()
            def positions_and_rectangles():
                positions = arrange_rectangles(area, [note.size() for note in ordered], alignment)
                return positions, [QRect(position, note.size()) for position, note in zip(positions, ordered)]
            positions, rectangles = positions_and_rectangles()
            for details_height in (160, 100, 80):
                overlaps = any(rectangle.intersects(other) for i, rectangle in enumerate(rectangles)
                               for other in rectangles[i + 1:])
                if not overlaps:
                    break
                for note in ordered:
                    if note.expanded:
                        note.details_scroll.setFixedHeight(details_height)
                        note.adjustSize()
                positions, rectangles = positions_and_rectangles()
            for note, position in zip(ordered, positions):
                if note.drag_offset is None:
                    note.move(position)
                    self.store.update_position(note.event.id, position.x(), position.y())
                    note.keep_below_windows()
        finally:
            self._arranging_desktop = False

    def sync_notes(self):
        if QApplication.instance().property("shutting_down"):
            return
        self.show_capture_notepad()
        try:
            self.schedule_midnight()
            self.today_label.setText("오늘 날짜: " + date.today().isoformat())
            due = self.store.reminder_events()
            font_size = int(self.store.get_setting("font_size", 12))
            if self.desktop_mode == "list":
                for note in list(self.notes.values()):
                    note.finish()
                self.notes.clear()
                if due:
                    if self.desktop_list is None:
                        self.desktop_list = DesktopListWindow(due, font_size)
                        self.desktop_list.completed.connect(self.complete_event)
                        self.desktop_list.moved.connect(self.save_list_position)
                        self.desktop_list.size_changed.connect(self.desktop_size_changed)
                        self.desktop_list.dismissed_by_user.connect(self.dismiss_desktop_list)
                    else:
                        self.desktop_list.update_events(due)
                    if not self._desktop_list_dismissed and not self.desktop_list.isVisible():
                        self.desktop_list.show()
                elif self.desktop_list is not None:
                    self.desktop_list.finish()
                    self.desktop_list = None
            else:
                if self.desktop_list is not None:
                    self.desktop_list.finish()
                    self.desktop_list = None
                ids = {int(event.id) for event in due}
                for id_ in list(self.notes):
                    if id_ not in ids:
                        self.notes.pop(id_).finish()
                for index, event in enumerate(due):
                    id_ = int(event.id)
                    if id_ in self._dismissed_note_ids:
                        continue
                    if id_ in self.notes:
                        self.notes[id_].update_event(event)
                        if not self.notes[id_].isVisible():
                            self.notes[id_].show()
                    else:
                        note = NoteWindow(event, font_size)
                        note.completed.connect(self.complete_event)
                        note.moved.connect(self.save_note_position)
                        note.size_changed.connect(self.desktop_size_changed)
                        note.dismissed_by_user.connect(lambda event_id=id_: self.dismiss_note(event_id))
                        self.notes[id_] = note
                        self.safe_position(note, event, index)
                        note.show()
            self.arrange_desktop_bottom()
            self.tray.setToolTip(f"{APP_NAME} · 미완료 알림 {len(due)}개")
            self.update_note_reason(self.agenda.currentItem())
            self.queue_daily_backup()
        except Exception:
            logging.exception("Reminder refresh failed")
            if hasattr(self, "status_label"):
                self.status_label.setText("일정을 불러오지 못했습니다. 데이터 폴더와 저장 공간을 확인해 주세요.")

    def refresh_startup_reminders(self):
        # Explorer and display work areas can settle after a login-started app.
        # Reuse the normal refresh so filters and tools closed by the user stay
        # respected while their desktop layer and positions are re-applied.
        if (QApplication.instance().property("shutting_down")
                or self._startup_refresh_remaining <= 0):
            self.startup_refresh_timer.stop()
            self._startup_refresh_remaining = 0
            return
        self._startup_refresh_remaining -= 1
        self.sync_notes()
        if self._startup_refresh_remaining == 0:
            self.startup_refresh_timer.stop()

    def queue_daily_backup(self):
        # Let the reminder windows paint before doing optional backup I/O.
        if (self.args.smoke_test or QApplication.instance().property("shutting_down")
                or self._daily_backup_pending or time.monotonic() < self._daily_backup_retry_at):
            return
        self._daily_backup_pending = True
        QTimer.singleShot(500, self.run_daily_backup)

    def run_daily_backup(self):
        try:
            if (self.args.smoke_test or QApplication.instance().property("shutting_down")
                    or time.monotonic() < self._daily_backup_retry_at):
                return
            today = date.today().isoformat()
            if self.store.get_setting("last_backup_date") == today:
                return
            self.store.backup()
            self.store.set_setting("last_backup_date", today)
            self._daily_backup_retry_at = 0.0
        except Exception:
            self._daily_backup_retry_at = time.monotonic() + 60
            logging.exception("Automatic daily backup failed; reminders remain visible")
            self.status_label.setText(
                "일정은 표시되어 있습니다. 자동 백업에 실패하여 1분 후 다시 시도합니다. "
                "데이터 폴더와 저장 공간을 확인해 주세요.")
        finally:
            self._daily_backup_pending = False

    def schedule_midnight(self):
        if not hasattr(self, "midnight_timer"):
            return
        now = datetime.now()
        next_midnight = datetime.combine(now.date() + timedelta(days=1), datetime.min.time())
        delay = max(1, int((next_midnight - now).total_seconds() * 1000))
        remaining = self.midnight_timer.remainingTime()
        if not self.midnight_timer.isActive() or abs(remaining - delay) > 2000:
            self.midnight_timer.start(delay)

    def on_midnight(self):
        self.refresh()
        self.schedule_midnight()

    def save_note_position(self, id_, x, y):
        try:
            self.store.set_setting("desktop_layout", "free")
            self.sync_alignment_controls("free")
            self.store.update_position(id_, x, y)
        except Exception:
            self.status_label.setText("메모 위치를 저장하지 못했습니다.")

    def complete_event(self, id_):
        try:
            self.store.complete_event(id_)
            self._dismissed_note_ids.discard(id_)
            if id_ in self.notes:
                self.notes.pop(id_).finish()
            self.refresh()
            self.status_label.setText("완료 처리했습니다. 완료 기록에서 되돌릴 수 있습니다.")
        except Exception as exc:
            self.show_error("완료 상태를 저장하지 못했습니다. 메모를 유지합니다.\n" + str(exc))

    def add_manual(self):
        dialog = EventDialog(parent=self)
        dialog.date_field.setText(self.calendar.selectedDate().toString("yyyy-MM-dd"))
        if dialog.exec_() == QDialog.Accepted:
            try:
                self.store.upsert_event(dialog.event)
                self.refresh()
                self.status_label.setText("일정을 저장했습니다.")
            except Exception as exc:
                self.show_error(str(exc))

    def edit_event(self, id_):
        if id_ is None:
            return
        event = self.store.get_event(id_)
        if event is None:
            return
        dialog = EventDialog(event, self)
        if dialog.exec_() == QDialog.Accepted:
            try:
                saved = self.store.upsert_event(dialog.event)
                selected = self.calendar.selectedDate().toString("yyyy-MM-dd")
                if not saved.start_date <= selected <= (saved.end_date or saved.start_date):
                    self.calendar.setSelectedDate(QDate.fromString(saved.start_date, "yyyy-MM-dd"))
                self.refresh()
                for row in range(self.agenda.count()):
                    item = self.agenda.item(row)
                    if item.data(Qt.UserRole) == saved.id:
                        self.agenda.setCurrentItem(item)
                        break
                self.status_label.setText("일정을 수정했습니다.")
            except Exception as exc:
                self.show_error(str(exc))

    def table_action(self, table, mode):
        row = table.currentRow()
        if row < 0 or table.item(row, 0) is None:
            return
        id_ = table.item(row, 0).data(Qt.UserRole)
        if mode == "review":
            self.edit_event(id_)
        else:
            try:
                self.store.reopen_event(id_)
                self.refresh()
            except Exception as exc:
                self.show_error(str(exc))

    def import_text(self):
        text = self.text_input.toPlainText().strip()
        if not text:
            self.import_files()
            return
        source = self.import_preview[1] if self.import_preview and text == self.import_preview[0] else "붙여넣은 텍스트"
        self.process_import(text, source)

    def import_pasted_text(self):
        self.import_preview = None
        if self.text_input.toPlainText().strip():
            self.import_text()

    def receive_import_text(self, text, source):
        self.import_preview = (text.strip(), source)
        self.import_preview_records = [(source, text)]
        self.text_input.setPlainText(text)
        self.process_import(text, source)

    def sync_recognized_text(self, text, source):
        """Keep edited recognition text as a draft without parsing schedules."""
        previous_text = self.text_input.toPlainText()
        self.import_preview = (text.strip(), source)
        # Once the combined text is edited, per-file snapshots no longer
        # describe it. A later explicit analysis must read the edited text.
        self.import_preview_records = []
        if self.text_input.toPlainText() != text:
            self.text_input.setPlainText(text)
        if self.capture_notepad.editor.toPlainText() in ("", previous_text):
            self.capture_notepad.set_draft(text)
        self.draft_timer.start(300)

    def process_import(self, text, source):
        try:
            self._process_import(text, source)
        except Exception as exc:
            self.show_error("일정을 분석하지 못했습니다.\n" + str(exc))

    def _process_import(self, text, source):
        records = getattr(self, "import_preview_records", []) if self.import_preview == (text.strip(), source) else []
        candidates = [event for name, raw in records for event in parse_schedule_text(raw, source=name)] if records else parse_schedule_text(text, source=source)
        self.queue_import_review(candidates, source, text)

    def queue_import_review(self, candidates, source, text, errors=None):
        if QApplication.instance().property("shutting_down"):
            return
        summary = f"분석한 일정 {len(candidates)}건 · 아직 등록되지 않았습니다."
        self.import_result.setPlainText(summary + "\n\n" + text[:18000])
        self.status_label.setText(source + ": " + summary)
        self.capture_notepad.status_label.setText("확인 창에서 내용을 수정한 뒤 ‘등록’을 눌러 주세요.")
        self.import_queue.append((candidates, source, text, errors or []))
        self.show_next_import_review()

    def show_next_import_review(self):
        if self.import_dialog or not self.import_queue or QApplication.instance().property("shutting_down"):
            return
        candidates, source, text, errors = self.import_queue.pop(0)
        review_candidates, existing_ids = [], set()
        for candidate in candidates:
            existing = self.store.find_duplicate(candidate)
            if existing is not None:
                if existing.id in existing_ids:
                    continue
                existing_ids.add(existing.id)
            review_candidates.append(existing or candidate)
        candidates = review_candidates
        dialog = ImportReviewDialog(candidates, errors=errors, recognized_text=text)
        dialog.setWindowIcon(self.icon)
        if not candidates:
            dialog.source_text.setPlainText(text)
        self.import_dialog = dialog
        dialog.recognized_text_changed.connect(lambda edited: self.sync_recognized_text(edited, source))
        dialog.confirmed.connect(lambda selected: self.confirm_import(
            dialog, selected, source, dialog.recognized_editor.toPlainText()))
        def finished(result):
            if self.import_dialog is dialog:
                self.import_dialog = None
            if result != QDialog.Accepted:
                self.status_label.setText("등록을 취소했습니다. 입력한 원문은 남아 있습니다.")
                self.capture_notepad.status_label.setText("등록을 취소했습니다. 초안을 수정해 다시 확인할 수 있습니다.")
            dialog.deleteLater()
            QTimer.singleShot(0, self.show_next_import_review)
        dialog.finished.connect(finished)
        dialog.show()
        dialog.raise_()
        dialog.activateWindow()

    def confirm_import(self, dialog, selected, source, text):
        try:
            candidates = [replace(event) for event in selected]
            added, updated, duplicates = self.store.save_reviewed_events(candidates)
        except Exception as exc:
            dialog.error_label.setText("일정을 저장하지 못했습니다. 다시 등록할 수 있습니다.\n" + str(exc))
            return
        summary = f"새 일정 {added}건 · 중복 {duplicates}건"
        if updated:
            summary = f"새 일정 {added}건 · 수정 {updated}건 · 중복 {duplicates}건"
        dialog.accept()
        self.refresh()
        matched = {event.id: event for candidate in candidates
                   if (event := (self.store.get_event(candidate.id) if candidate.id is not None else
                                 self.store.find_duplicate(candidate))) is not None}
        if self.desktop_mode == "list":
            rendered = set(self.desktop_list.events) if self.desktop_list and self.desktop_list.isVisible() else set()
        else:
            rendered = {id_ for id_, note in self.notes.items() if note.isVisible()}
        shown = sum(id_ in rendered for id_ in matched)
        eligible_ids = {event.id for event in self.store.reminder_events()}
        hidden_by_range = sum(event.status == "open" and id_ not in eligible_ids
                              for id_, event in matched.items())
        hidden_by_close = sum(event.status == "open" and id_ in eligible_ids and id_ not in rendered
                              for id_, event in matched.items())
        completed = sum(event.status == "done" for event in matched.values())
        message = (f"이미 등록된 일정 {duplicates}건입니다. 기존 일정을 유지했습니다." if not added and not updated and duplicates else
                   f"새 일정 {added}건을 등록했습니다." + (f" 이미 등록된 {duplicates}건은 중복으로 추가하지 않았습니다." if duplicates else ""))
        if updated:
            message += f" 기존 일정 {updated}건을 수정했습니다."
        if shown:
            message += f" 바탕화면에 {shown}건을 표시했습니다."
        if hidden_by_close:
            message += f" {hidden_by_close}건은 바탕화면 도구를 닫아 숨겨져 있습니다. 트레이의 '바탕화면 알림 다시 열기'에서 다시 표시하세요."
        if hidden_by_range:
            message += f" {hidden_by_range}건은 설정한 표시 기간 밖입니다. 설정에서 기간을 늘리거나 ‘날짜 제한 없이’ 표시를 선택하세요."
        if completed:
            message += f" 완료된 기존 일정 {completed}건은 완료 상태를 유지합니다. 다시 표시하려면 완료 기록에서 완료 취소하세요."
        self.import_result.setPlainText(summary + "\n" + message + "\n\n" + text[:18000])
        self.status_label.setText(source + ": " + message)
        self.capture_notepad.status_label.setText(message)

    def fill_sample(self):
        today = date.today()
        weekday = "월화수목금토일"[today.weekday()]
        self.text_input.setPlainText(f"{today.year}년 {today.month}월 {today.day}일({weekday}) 15:30 교직원 협의회, 장소: 2층 회의실")

    def paste_clipboard(self):
        self.text_input.accept_mime(QApplication.clipboard().mimeData())

    def receive_import_payload(self, payload, origin="desktop"):
        kind = payload.get("kind")
        if kind == "files":
            self.import_file_paths(payload["paths"], origin)
        elif kind == "image":
            image = qimage_to_pil(payload["image"])
            def received(text):
                if origin == "desktop":
                    self.capture_notepad.append_recognized_text(text)
                    self.save_capture_draft()
                self.receive_import_text(text, "클립보드 / 드래그 이미지")
            self.run_task(lambda: self.ocr.recognize(image), received)
        elif kind == "text":
            if origin == "desktop":
                self.save_capture_draft()
            self.receive_import_text(payload["text"], "붙여넣은 텍스트")

    def import_files(self):
        extensions = " ".join("*" + suffix for suffix in SUPPORTED_EXTENSIONS)
        paths, _ = QFileDialog.getOpenFileNames(self, "일정 파일 가져오기", "", f"지원 문서 ({extensions});;모든 파일 (*)")
        if not paths:
            return
        self.import_file_paths(paths)

    def import_file_paths(self, paths, origin="manager"):
        def read_files():
            results = []
            for path in paths:
                try:
                    results.append((Path(path).name, extract_file_text(path, self.ocr.recognize), None))
                except Exception as exc:
                    results.append((Path(path).name, "", str(exc)))
            return results
        def receive(results):
            errors, candidates, parts, names = [], [], [], []
            for source, text, error in results:
                if error:
                    errors.append(source + ": " + error)
                else:
                    candidates.extend(parse_schedule_text(text, source=source))
                    parts.append(text)
                    names.append(source)
            text = "\n\n".join(parts)
            source = ", ".join(names) or "가져온 파일"
            self.import_preview = (text.strip(), source)
            self.import_preview_records = [(name, raw) for name, raw, error in results if not error]
            self.text_input.setPlainText(text)
            if origin == "desktop":
                self.capture_notepad.append_recognized_text(text)
                self.save_capture_draft()
            self.queue_import_review(candidates, source, text, errors)
        self.run_task(read_files, receive)

    def capture_screen(self):
        if not self.ocr.available():
            self.show_error(self.ocr.diagnostics())
            return
        self.hide()
        QTimer.singleShot(220, self.begin_capture)

    def begin_capture(self):
        screen = QApplication.screenAt(QCursor.pos()) or QApplication.primaryScreen()
        self.overlay = CaptureOverlay(screen)
        self.overlay.captured.connect(self.receive_capture)
        self.overlay.canceled.connect(self.show_manager)
        self.overlay.show()
        self.overlay.activateWindow()

    def receive_capture(self, image):
        self.show_manager()
        self.navigation.setCurrentRow(1)
        self.run_task(lambda: self.ocr.recognize(image), lambda text: self.receive_import_text(text, "화면 캡처"))

    def run_task(self, function, callback):
        task = Task(function)
        self.jobs.add(task)
        self.capture_notepad.set_busy(True)
        self.status_label.setText("PC 안에서 자료를 분석하고 있습니다…")
        self.analyze_button.setEnabled(False)
        self.capture_button.setEnabled(False)
        self.file_button.setEnabled(False)
        def deliver(result):
            if QApplication.instance().property("shutting_down"):
                return
            try:
                callback(result)
            except Exception as exc:
                self.show_error("자료 처리 중 문제가 발생했습니다.\n" + str(exc))
        task.signals.result.connect(deliver)
        task.signals.error.connect(self.show_error)
        def finish():
            self.jobs.discard(task)
            self.analyze_button.setEnabled(not self.jobs)
            self.capture_button.setEnabled(not self.jobs)
            self.file_button.setEnabled(not self.jobs)
            self.capture_notepad.set_busy(bool(self.jobs))
            if not self.jobs:
                self.capture_notepad.status_label.setText(self.status_label.text())
        task.signals.finished.connect(finish)
        self.pool.start(task)

    def backup_data(self):
        path, _ = QFileDialog.getSaveFileName(self, "백업 저장", "학교알리미-백업.sqlite3", "SQLite 백업 (*.sqlite3)")
        if path:
            try:
                self.store.backup(path)
                self.status_label.setText("백업 파일을 저장했습니다.")
            except Exception as exc:
                self.show_error(str(exc))

    def reset_schedule_range(self):
        self._reset_schedules(self.reset_from_date.date().toString("yyyy-MM-dd"),
                              self.reset_to_date.date().toString("yyyy-MM-dd"))

    def reset_all_schedules(self):
        self._reset_schedules()

    def _reset_schedules(self, start=None, end=None):
        if self.jobs:
            QMessageBox.information(self, "자료 분석 중", "자료 분석이 끝난 뒤 일정을 초기화해 주세요.")
            return
        try:
            event_ids = self.store.schedule_reset_ids(start, end)
            if not event_ids:
                QMessageBox.information(self, "일정 초기화", "초기화할 일정이 없습니다.")
                return
            scope = f"{start} ~ {end} 기간과 겹치는 일정" if start is not None else "전체 일정 (날짜 없는 일정 포함)"
            answer = QMessageBox.warning(
                self, "일정 초기화 확인",
                f"{scope} {len(event_ids)}건을 초기화할까요?\n"
                "미완료·완료·확인 필요 일정이 삭제됩니다. 삭제 전에 자동 백업을 만들며 백업에서 복원할 수 있습니다.\n"
                "표시 기간 설정과 메모 초안은 유지됩니다.",
                QMessageBox.Ok | QMessageBox.Cancel, QMessageBox.Cancel,
            )
            if answer != QMessageBox.Ok:
                return
            self.save_capture_draft()
            deleted, backup_path = self.store.reset_events(event_ids)
            self.import_queue.clear()
            if self.import_dialog is not None:
                self.import_dialog.reject()
            self._dismissed_note_ids.difference_update(event_ids)
            self.refresh()
            message = f"일정 {deleted}건을 초기화했습니다."
            if backup_path is not None:
                message += "\n자동 백업: " + str(backup_path)
            self.reset_result_label.setText(message)
            self.status_label.setText(message)
        except Exception as exc:
            self.show_error("일정을 초기화하지 못했습니다.\n" + str(exc))

    def restore_data(self):
        path, _ = QFileDialog.getOpenFileName(self, "백업 복원", "", "SQLite 백업 (*.sqlite3)")
        if not path:
            return
        if QMessageBox.question(self, "백업 복원", "현재 일정을 먼저 자동 백업한 뒤 선택한 백업으로 복원합니다. 계속할까요?") != QMessageBox.Yes:
            return
        try:
            self.store.restore(path)
            self.import_queue.clear()
            if self.import_dialog is not None:
                self.import_dialog.reject()
            self.sync_reminder_controls()
            self.refresh()
            self.status_label.setText("백업을 복원했습니다.")
        except Exception as exc:
            self.show_error(str(exc))

    def show_error(self, message):
        self.status_label.setText(message[:180])
        QMessageBox.warning(self, "학교 알리미", message)

    def on_shutdown(self):
        QApplication.instance().setProperty("shutting_down", True)
        self.startup_refresh_timer.stop()
        self._startup_refresh_remaining = 0
        self.draft_timer.stop()
        self.save_capture_draft()
        icon_position = self.capture_icon_position or self.capture_notepad.pos()
        self.save_capture_position(icon_position.x(), icon_position.y())
        # Windows can cancel a shutdown, and commitDataRequest/aboutToQuit may
        # both call this method. Keep the widget alive for either path.
        self.capture_notepad.allow_close = True
        self.capture_notepad.close()
        self.import_queue.clear()
        if self.import_dialog:
            self.import_dialog.reject()
        for id_, note in self.notes.items():
            try:
                self.store.update_position(id_, note.x(), note.y())
            except Exception:
                pass
        if self.desktop_list:
            self.store.set_setting("list_x", self.desktop_list.x())
            self.store.set_setting("list_y", self.desktop_list.y())
        self.tray.hide()

    def on_session_end(self, _manager):
        # Allow Windows shutdown/logoff even while reminders are unfinished.
        self.on_shutdown()
        def recover_canceled_shutdown():
            app = QApplication.instance()
            app.setProperty("shutting_down", False)
            if not self.args.smoke_test:
                self.tray.show()
            self.capture_notepad.allow_close = False
            self.show_capture_notepad()
            self.sync_notes()
            for note in self.notes.values():
                if not note.isVisible() and not note.dismissed:
                    note.show()
                    note.keep_below_windows()
            if self.desktop_list and not self.desktop_list.isVisible() and not self._desktop_list_dismissed:
                self.desktop_list.show()
                self.desktop_list.keep_below_windows()
        QTimer.singleShot(10_000, recover_canceled_shutdown)


def smoke_test(app, window, store, report_path: Path):
    report = {"success": False, "version": APP_VERSION, "checks": {}}
    def close_with_answer(widget, answer, button=None, activate=None):
        def respond():
            box = app.activeModalWidget()
            if isinstance(box, QMessageBox):
                warning_path = report_path.with_name(report_path.stem + "-close-warning.png")
                if not warning_path.exists():
                    box.grab().save(str(warning_path))
                box.button(answer).click()
        QTimer.singleShot(0, respond)
        (activate or (button.click if button else widget.close))()
        app.processEvents()
    try:
        event = store.upsert_event(Event(title="실행 검증 일정", start_date=date.today().isoformat(),
                                        time="15:30", place="테스트 회의실", source="실행 검증"))
        window.refresh()
        app.processEvents()
        note = window.notes[int(event.id)]
        report["checks"]["independent_window"] = note.parent() is None
        report["checks"]["bottom_layer_flags"] = bool(note.windowFlags() & Qt.WindowStaysOnBottomHint) and not bool(note.windowFlags() & Qt.WindowStaysOnTopHint)
        report["checks"]["compact_sticky_title_only"] = not note.expanded and not note.details_scroll.isVisible() and not note.complete_button.isVisible()
        report["checks"]["deadline_date_and_weekday_left_of_title"] = (
            note.date_label.text() == format_event_deadline(event) and
            note.date_label.geometry().right() < note.title_label.geometry().left() and
            "월화수목금토일"[date.today().weekday()] in note.date_label.text())
        report["checks"]["today_reminder_is_orange"] = note.property("urgency") == "due" and note.grab().toImage().pixelColor(3, 3).name() == event_urgency(event).background
        note.title_label.click()
        app.processEvents()
        report["checks"]["title_click_reveals_full_details"] = note.expanded and note.details_scroll.isVisible() and note.complete_button.isVisible()
        note.title_label.click()
        app.processEvents()
        report["checks"]["title_click_collapses_details"] = not note.expanded and not note.details_scroll.isVisible()
        window.show()
        app.processEvents()
        window.grab().save(str(report_path.with_name(report_path.stem + "-manager.png")))
        note.grab().save(str(report_path.with_name(report_path.stem + "-note.png")))
        close_with_answer(window, QMessageBox.Cancel, window.close_button)
        report["checks"]["manager_close_cancel_keeps_visible"] = window.isVisible()
        close_with_answer(window, QMessageBox.Ok, window.close_button)
        report["checks"]["manager_close_keeps_note"] = not window.isVisible() and note.isVisible()
        capture = window.capture_notepad
        report["checks"]["always_visible_desktop_notepad"] = capture.parent() is None and capture.isVisible() and not capture.editor.isReadOnly()
        report["checks"]["desktop_notepad_starts_as_icon"] = capture.collapsed and capture.icon_button.isVisible() and not capture.editor.isVisible() and not capture.isMinimized()
        report["checks"]["notepad_icon_is_not_cropped"] = (
            capture.width() == 72 and capture.height() == 70 and
            capture.icon_button.height() >= capture.icon_button.iconSize().height())
        report["checks"]["notepad_icon_has_no_close_button"] = (
            capture.compact_panel.findChildren(QPushButton) == [capture.icon_button])
        report["checks"]["notepad_always_on_top"] = (
            bool(capture.windowHandle().flags() & Qt.WindowStaysOnTopHint) and
            not bool(capture.windowHandle().flags() & Qt.WindowStaysOnBottomHint))
        icon_image = capture.grab().toImage()
        report["checks"]["notepad_icon_has_transparent_background"] = (
            capture.testAttribute(Qt.WA_TranslucentBackground) and
            all(icon_image.pixelColor(x, y).alpha() == 0
                for x, y in ((0, 0), (icon_image.width() - 1, 0),
                             (0, icon_image.height() - 1),
                             (icon_image.width() - 1, icon_image.height() - 1))))
        report["checks"]["notepad_icon_has_no_visible_text"] = not any(
            label.isVisibleTo(capture) and label.text().strip()
            for label in capture.findChildren(QLabel))
        icon_position = QPoint(capture.pos())
        capture.grab().save(str(report_path.with_name(report_path.stem + "-notepad-icon.png")))
        from PyQt5.QtTest import QTest
        from PyQt5.QtGui import QContextMenuEvent
        def open_icon_menu():
            point = QPoint(36, 35)
            QApplication.sendEvent(capture.icon_button, QContextMenuEvent(
                QContextMenuEvent.Mouse, point, capture.icon_button.mapToGlobal(point)))
            app.processEvents()
        open_icon_menu()
        report["checks"]["notepad_icon_right_click_opens_close_menu"] = (
            capture.icon_menu.isVisible() and capture.icon_close_action.text() == "닫기")
        capture.icon_menu.grab().save(str(report_path.with_name(report_path.stem + "-notepad-menu.png")))
        def click_menu_close():
            QTest.mouseClick(capture.icon_menu, Qt.LeftButton,
                             pos=capture.icon_menu.actionGeometry(capture.icon_close_action).center())
        close_with_answer(capture, QMessageBox.Cancel, activate=click_menu_close)
        report["checks"]["icon_menu_close_cancel_preserves_icon"] = (
            capture.isVisible() and capture.collapsed and not capture.dismissed)
        open_icon_menu()
        close_with_answer(capture, QMessageBox.Ok, activate=click_menu_close)
        window.sync_notes()
        report["checks"]["icon_menu_close_confirmation_stays_closed"] = (
            capture.dismissed and not capture.isVisible())
        window.show_capture_notepad(expand=True)
        capture.set_collapsed(True)
        app.processEvents()
        QTest.mouseDClick(capture.icon_button, Qt.LeftButton)
        app.processEvents()
        report["checks"]["icon_double_click_opens_editor"] = not capture.collapsed and capture.editor.isVisible() and capture.minimize_button.isVisible()
        report["checks"]["expanded_notepad_keeps_opaque_background"] = (
            capture.grab().toImage().pixelColor(3, 3).alpha() == 255)
        capture.editor.setPlainText("메모장 아이콘 전환 검증 초안")
        QTest.mouseDClick(capture.title_label, Qt.LeftButton)
        app.processEvents()
        report["checks"]["header_double_click_collapses_without_losing_draft"] = capture.collapsed and capture.icon_button.isVisible() and capture.editor.toPlainText() == "메모장 아이콘 전환 검증 초안" and capture.pos() == icon_position
        window.sync_notes()
        app.processEvents()
        report["checks"]["periodic_refresh_keeps_notepad_collapsed"] = capture.collapsed
        window.show_capture_notepad(expand=True)
        app.processEvents()
        capture.minimize_button.click()
        app.processEvents()
        report["checks"]["minimize_button_keeps_desktop_icon_visible"] = capture.collapsed and capture.isVisible() and not capture.isMinimized()
        capture.expand_notepad()
        app.processEvents()
        close_with_answer(capture, QMessageBox.Cancel, capture.close_button)
        report["checks"]["notepad_close_cancel_keeps_visible"] = capture.isVisible()
        capture.grab().save(str(report_path.with_name(report_path.stem + "-notepad.png")))
        close_with_answer(capture, QMessageBox.Ok, capture.close_button)
        window.sync_notes()
        report["checks"]["notepad_close_preserves_draft_and_stays_closed"] = (
            not capture.isVisible() and capture.dismissed and
            store.get_setting("capture_draft") == capture.editor.toPlainText())
        window.show_capture_notepad(expand=True)
        report["checks"]["notepad_reopens_from_tray_action"] = capture.isVisible() and not capture.dismissed
        close_with_answer(note, QMessageBox.Cancel, note.close_button)
        report["checks"]["note_close_cancel_keeps_visible"] = note.isVisible()
        close_with_answer(note, QMessageBox.Ok, note.close_button)
        window.sync_notes()
        report["checks"]["note_close_preserves_unfinished_and_stays_closed"] = (
            not note.isVisible() and store.get_event(event.id).status == "open")
        window.restore_desktop_tools()
        report["checks"]["note_reopens_from_tray_action"] = note.isVisible() and not note.dismissed
        note.title_label.click()
        note.complete_button.click()
        app.processEvents()
        report["checks"]["complete_saved"] = store.get_event(event.id).status == "done"
        report["checks"]["complete_removes_note"] = int(event.id) not in window.notes
        store.reopen_event(event.id)
        window.sync_notes()
        report["checks"]["reopen_restores_note"] = int(event.id) in window.notes
        report["checks"]["offline_ocr_ready"] = window.ocr.available()
        sample = window.ocr.resources_path / "verification-sample.png"
        text = window.ocr.recognize(Image.open(sample).convert("RGB"))
        report["checks"]["korean_ocr_inference"] = "교직원" in text and "비상연락망" in text
        candidates = parse_schedule_text(text, source="한국어 OCR 실행 검증")
        report["checks"]["ocr_schedule_extraction"] = len(candidates) >= 2 and any(e.start_date == "2026-10-15" for e in candidates)
        from docx import Document
        import openpyxl
        fixture = window.args.data_dir / "document-check.docx"
        doc = Document()
        doc.add_paragraph("2026년 10월 15일(목) 15:00 교직원 회의")
        doc.save(fixture)
        report["checks"]["docx_import"] = "교직원 회의" in extract_file_text(fixture)
        workbook = openpyxl.Workbook()
        workbook.active.append(["2026-10-15", "교직원 회의"])
        xlsx = window.args.data_dir / "document-check.xlsx"
        workbook.save(xlsx)
        workbook.close()
        report["checks"]["xlsx_import"] = "교직원 회의" in extract_file_text(xlsx)
        import pdfplumber, pypdfium2, olefile
        report["checks"]["pdf_hwp_modules"] = all((pdfplumber, pypdfium2, olefile))
        mime = QMimeData()
        mime.setText(date.today().isoformat() + " 16:00 붙여넣기 검증 회의")
        window.text_input.clear()
        window.text_input.insertFromMimeData(mime)
        loop = QEventLoop()
        QTimer.singleShot(250, loop.quit)
        loop.exec_()
        report["checks"]["paste_requires_final_confirmation"] = window.import_dialog is not None and not any("붙여넣기 검증" in e.title for e in store.list_events())
        recognized_dialog = window.import_dialog
        recognized_before = recognized_dialog.selected_events()
        corrected_body = recognized_dialog.recognized_editor.toPlainText() + "\n안내: 일정 외 인식 텍스트도 수정할 수 있습니다."
        recognized_dialog.recognized_editor.setPlainText(corrected_body)
        report["checks"]["recognized_body_edit_leaves_schedule_fields_unchanged"] = (
            recognized_dialog.selected_events() == recognized_before and window.text_input.toPlainText() == corrected_body)
        window.save_capture_draft()
        report["checks"]["recognized_body_edit_saved_locally"] = store.get_setting("recognized_text_draft") == corrected_body
        window.import_dialog.table.item(0, ImportReviewDialog.TITLE).setText("수정한 붙여넣기 검증 회의")
        window.import_dialog.grab().save(str(report_path.with_name(report_path.stem + "-confirmation.png")))
        window.import_dialog.register_button.click()
        app.processEvents()
        report["checks"]["paste_edits_saved_only_after_register"] = any(e.title == "수정한 붙여넣기 검증 회의" for e in store.list_events())
        report["checks"]["registered_result_keeps_corrected_recognized_body"] = corrected_body in window.import_result.toPlainText()
        capture.editor.setPlainText("2099년 10월 15일 취소 검증 회의")
        window.review_capture_draft()
        report["checks"]["typed_draft_opens_confirmation"] = window.import_dialog is not None and not any("취소 검증" in e.title for e in store.list_events())
        window.import_dialog.reject()
        app.processEvents()
        window.save_capture_draft()
        report["checks"]["cancel_preserves_local_draft_without_registration"] = "취소 검증" in store.get_setting("capture_draft", "") and not any("취소 검증" in e.title for e in store.list_events())
        original_deadline = "10월 30일까지 수행평가 제출하세요"
        capture.editor.setPlainText(original_deadline)
        window.review_capture_draft()
        report["checks"]["confirmation_preserves_original_deadline_sentence"] = (
            window.import_dialog.table.item(0, ImportReviewDialog.TITLE).text() == original_deadline)
        window.import_dialog.register_button.click()
        app.processEvents()
        deadline_event = next(e for e in store.list_events() if e.title == original_deadline)
        report["checks"]["registered_sticky_preserves_original_deadline_sentence"] = (
            deadline_event.start_date == f"{date.today().year}-10-30" and
            window.notes[int(deadline_event.id)].title_label.full_title == original_deadline and
            original_deadline in deadline_event.source_text)
        future_date = (date.today() + timedelta(days=2)).isoformat()
        future = store.upsert_event(Event(title="미래 일정 검증", start_date=future_date, status="open"))
        review = store.upsert_event(Event(title="확인 필요 일정 검증", start_date=date.today().isoformat(),
                                         status="review", notes="확인 필요: 연도를 확인해 주세요."))
        window.refresh()
        report["checks"]["two_day_deadline_already_displayed"] = int(future.id) in window.notes and window.notes[int(future.id)].property("urgency") == "two_days"
        report["checks"]["review_today_not_displayed"] = int(review.id) not in window.notes
        window.calendar.setSelectedDate(QDate.fromString(future_date, "yyyy-MM-dd"))
        report["checks"]["future_display_date_explained"] = "바탕화면에 표시" in window.note_reason.text() and "D-2" in window.note_reason.text()
        ahead = store.upsert_event(Event(title="마감 3일 전 검증", start_date=(date.today() + timedelta(days=3)).isoformat()))
        later = store.upsert_event(Event(title="마감 4일 전 검증", start_date=(date.today() + timedelta(days=4)).isoformat()))
        window.refresh()
        report["checks"]["default_unlimited_shows_future_reminders_immediately"] = window.reminder_unlimited.isChecked() and int(ahead.id) in window.notes and int(later.id) in window.notes
        window.reminder_unlimited.setChecked(False)
        window.reminder_past_days.setValue(0)
        window.reminder_future_days.setValue(3)
        app.processEvents()
        report["checks"]["limited_future_day_boundary"] = int(ahead.id) in window.notes and int(later.id) not in window.notes and window.reminder_future_days.isEnabled()
        report["checks"]["range_filter_preserves_unfinished_status"] = store.get_event(later.id).status == "open"
        window.navigation.setCurrentRow(4)
        window.show()
        app.processEvents()
        window.grab().save(str(report_path.with_name(report_path.stem + "-range-settings.png")))
        window.reminder_future_days.setValue(0)
        app.processEvents()
        report["checks"]["today_only_excludes_future_reminders"] = int(ahead.id) not in window.notes and int(future.id) not in window.notes
        window.reminder_unlimited.setChecked(True)
        app.processEvents()
        report["checks"]["unlimited_restores_future_and_disables_day_controls"] = int(ahead.id) in window.notes and int(later.id) in window.notes and not window.reminder_past_days.isEnabled() and not window.reminder_future_days.isEnabled()
        window.navigation.setCurrentRow(0)
        ahead_note = window.notes[int(ahead.id)]
        stages = []
        for step in range(5):
            reference = date.today() + timedelta(days=step)
            ahead_note.update_event(ahead, today=reference)
            app.processEvents()
            stages.append(ahead_note.property("urgency"))
            report["checks"]["deadline_color_pixel_day_" + str(step)] = ahead_note.grab().toImage().pixelColor(3, 3).name() == event_urgency(ahead, reference).background
            ahead_note.grab().save(str(report_path.with_name(report_path.stem + "-urgency-" + str(step) + ".png")))
        report["checks"]["three_days_to_overdue_gradual_colors"] = stages == ["three_days", "two_days", "one_day", "due", "due"]
        window.sync_notes()
        report["checks"]["urgency_refresh_uses_current_date"] = ahead_note.property("urgency") == "three_days"
        store.upsert_event(replace(review, status="open"))
        window.refresh()
        report["checks"]["confirmed_today_displayed"] = int(review.id) in window.notes
        window.test_postit_button.click()
        test_events = [e for e in store.list_events() if e.source == TEST_NOTE_SOURCE]
        report["checks"]["test_button_displays_postit"] = len(test_events) == 1 and int(test_events[0].id) in window.notes
        test_id = int(test_events[0].id)
        window.test_postit_button.click()
        report["checks"]["test_button_no_duplicate"] = len([e for e in store.list_events() if e.source == TEST_NOTE_SOURCE]) == 1
        window.complete_event(test_id)
        window.test_postit_button.click()
        report["checks"]["completed_test_can_be_shown_again"] = test_id in window.notes and store.get_event(test_id).status == "open"
        monthly = parse_schedule_text("2026학년도 평가 계획\n9월, 10월 국어 평가")
        report["checks"]["month_only_auto_registration"] = (
            [e.start_date for e in monthly] == ["2026-09-01", "2026-10-01"] and
            all(e.status == "open" and e.date_precision == "month" and "확인 필요" not in e.notes for e in monthly))
        report["checks"]["precise_midnight_timer"] = (
            window.midnight_timer.isSingleShot() and window.midnight_timer.isActive() and
            window.midnight_timer.timerType() == Qt.PreciseTimer and
            0 < window.midnight_timer.remainingTime() <= 86_400_000)
        if window.args.smoke_document:
            from collections import Counter
            document_text = extract_file_text(window.args.smoke_document, window.ocr.recognize)
            plan = parse_schedule_text(document_text, source=window.args.smoke_document.name)
            report["checks"]["monthly_pdf_automatically_ready"] = bool(plan) and all(
                e.status == "open" and e.start_date.endswith("-01") and "확인 필요" not in e.notes for e in plan)
            report["checks"]["monthly_dates_cover_whole_month"] = bool(plan) and all(
                e.end_date and e.end_date.startswith(e.start_date[:7]) and e.end_date >= e.start_date for e in plan)
            report["document_summary"] = dict(Counter(e.start_date for e in plan))
            validation_store = Store(window.args.data_dir / "monthly-plan.sqlite3")
            try:
                added, duplicates = validation_store.import_candidates(plan)
                report["checks"]["monthly_pdf_rows_not_lost"] = added == len(plan) and duplicates == 0
                september = [e for e in plan if e.start_date == "2026-09-01"]
                report["checks"]["month_first_day_boundary"] = (
                    not validation_store.due_events("2026-08-31") and
                    len(validation_store.due_events("2026-09-01")) == len(september))
            finally:
                validation_store.close()
            # Exercise the actual empty-input button, file picker and background
            # worker instead of calling the extraction/registration functions alone.
            original_picker = QFileDialog.getOpenFileNames
            original_error = window.show_error
            picker_calls = []
            import_errors = []
            def select_document(*_args, **_kwargs):
                picker_calls.append(True)
                return [str(window.args.smoke_document)], ""
            try:
                QFileDialog.getOpenFileNames = select_document
                window.show_error = lambda message: import_errors.append(message)
                window.text_input.clear()
                window.analyze_button.click()
                wait_loop = QEventLoop()
                completion = QTimer()
                completion.timeout.connect(lambda: wait_loop.quit() if not window.jobs else None)
                completion.start(20)
                deadline = QTimer()
                deadline.setSingleShot(True)
                deadline.timeout.connect(wait_loop.quit)
                deadline.start(30_000)
                wait_loop.exec_()
                completion.stop()
                deadline.stop()
                report["checks"]["file_worker_requires_confirmation"] = window.import_dialog is not None and not any(e.source == window.args.smoke_document.name for e in store.list_events())
                window.import_dialog.register_button.click()
                app.processEvents()
                imported = [e for e in store.list_events() if e.source == window.args.smoke_document.name]
                report["checks"]["empty_registration_button_selects_file"] = len(picker_calls) == 1
                report["checks"]["pdf_file_worker_registration"] = not window.jobs and not import_errors and len(imported) == len(plan)
                report["checks"]["pdf_original_text_preview"] = window.text_input.toPlainText() == document_text
                window.navigation.setCurrentRow(1)
                window.show()
                app.processEvents()
                window.grab().save(str(report_path.with_name(report_path.stem + "-import.png")))
                window.analyze_button.click()
                window.import_dialog.register_button.click()
                app.processEvents()
                report["checks"]["pdf_preview_retry_preserves_source_and_avoids_duplicates"] = (
                    not import_errors and len(picker_calls) == 1 and
                    len([e for e in store.list_events() if e.source == window.args.smoke_document.name]) == len(plan) and
                    f"중복 {len(plan)}건" in window.import_result.toPlainText())
            finally:
                QFileDialog.getOpenFileNames = original_picker
                window.show_error = original_error
            report["checks"]["monthly_import_has_no_confirmation_message"] = "확인 필요" not in window.import_result.toPlainText()
            display_month = plan[0].start_date[:7]
            window.calendar.setSelectedDate(QDate.fromString(display_month + "-15", "yyyy-MM-dd"))
            displayed = [store.get_event(window.agenda.item(row).data(Qt.UserRole))
                         for row in range(window.agenda.count())
                         if window.agenda.item(row).data(Qt.UserRole) is not None]
            report["checks"]["monthly_calendar_midmonth_visible"] = (
                sum(e.source == window.args.smoke_document.name for e in displayed) ==
                sum(e.start_date.startswith(display_month) for e in plan))
            first = next(e for e in store.reminder_events() if e.source == window.args.smoke_document.name)
            report["checks"]["monthly_pdf_postit_visible"] = int(first.id) in window.notes and window.notes[int(first.id)].isVisible()
            report["checks"]["monthly_pdf_displays_month_without_day"] = first.date_precision == "month" and first.start_date not in window.notes[int(first.id)].details_label.text() and format_event_date(first) in window.notes[int(first.id)].details_label.text()
            report["checks"]["monthly_deadline_label_keeps_month_precision"] = window.notes[int(first.id)].date_label.text() == format_event_date(first)
            report["checks"]["monthly_reminders_keep_neutral_yellow"] = window.notes[int(first.id)].urgency_state.code == "month" and window.notes[int(first.id)].grab().toImage().pixelColor(3, 3).name() == event_urgency(first).background
            window.notes[int(first.id)].grab().save(str(report_path.with_name(report_path.stem + "-monthly-note.png")))
            window.notes[int(first.id)].title_label.click()
            app.processEvents()
            window.notes[int(first.id)].grab().save(str(report_path.with_name(report_path.stem + "-monthly-detail.png")))
            window.notes[int(first.id)].title_label.click()
        remaining_ids = {e.id for e in store.reminder_events()}
        window.mode_combo.setCurrentIndex(1)
        app.processEvents()
        listing = window.desktop_list
        report["checks"]["list_mode_is_independent_and_bottom"] = bool(listing and listing.parent() is None and listing.isVisible() and listing.windowFlags() & Qt.WindowStaysOnBottomHint and not window.notes)
        report["checks"]["list_mode_dates_and_titles"] = bool(
            listing and not listing.details_panel.isVisible() and
            [(listing.list_widget.item(i).data(DEADLINE_ROLE), listing.list_widget.item(i).text())
             for i in range(listing.list_widget.count())] ==
            [(format_event_deadline(e), e.title) for e in store.reminder_events()])
        report["checks"]["list_mode_scrollbar"] = bool(listing and listing.list_widget.verticalScrollBarPolicy() != Qt.ScrollBarAlwaysOff)
        if listing:
            today_row = next(i for i in range(listing.list_widget.count()) if listing.list_widget.item(i).data(Qt.UserRole) == event.id)
            listing.list_widget.scrollToItem(listing.list_widget.item(today_row))
            app.processEvents()
            row_rect = listing.list_widget.visualItemRect(listing.list_widget.item(today_row))
            report["checks"]["list_urgency_color_actually_painted"] = listing.list_widget.viewport().grab().toImage().pixelColor(row_rect.right() - 5, row_rect.center().y()).name() == event_urgency(event).background
            listing.grab().save(str(report_path.with_name(report_path.stem + "-list.png")))
            listing.set_selected_event(listing.list_widget.item(0).data(Qt.UserRole))
            app.processEvents()
            report["checks"]["list_click_reveals_details"] = listing.details_panel.isVisible() and listing.complete_button.isVisible()
            divider_image = listing.details_divider.grab().toImage()
            report["checks"]["list_details_have_dark_divider"] = (
                listing.details_divider.isVisible() and divider_image.height() >= 3 and
                divider_image.pixelColor(divider_image.width() // 2,
                                         divider_image.height() // 2).name() == "#493e1e")
            window.sync_notes()
            report["checks"]["list_details_survive_refresh"] = listing.expanded and listing.selected_event_id in remaining_ids
            listing.grab().save(str(report_path.with_name(report_path.stem + "-list-detail.png")))
            close_with_answer(listing, QMessageBox.Cancel, listing.close_button)
            report["checks"]["list_close_cancel_keeps_visible"] = listing.isVisible()
            close_with_answer(listing, QMessageBox.Ok, listing.close_button)
            window.sync_notes()
            report["checks"]["list_close_stays_closed_after_refresh"] = not listing.isVisible()
            window.restore_desktop_tools()
            report["checks"]["list_reopens_from_tray_action"] = listing.isVisible()
        window.mode_combo.setCurrentIndex(0)
        app.processEvents()
        report["checks"]["mode_switch_preserves_all_events"] = window.desktop_list is None and set(window.notes) == remaining_ids
        original_capture_position = window.capture_notepad.pos()
        first_note = window.notes[int(store.reminder_events()[0].id)]
        area = QApplication.primaryScreen().availableGeometry()
        for alignment in ("bottom", "top", "right", "left"):
            window.alignment_combo.setCurrentIndex(window.alignment_combo.findData(alignment))
            app.processEvents()
            edge = first_note.geometry()
            report.setdefault("alignment_geometry", {})[alignment] = {
                "note": [edge.x(), edge.y(), edge.width(), edge.height()],
                "area": [area.x(), area.y(), area.width(), area.height()]}
            correct_edge = (edge.bottom() == area.bottom() - 12 if alignment == "bottom" else
                            edge.top() == area.top() + 12 if alignment == "top" else
                            edge.right() == area.right() - 12 if alignment == "right" else
                            edge.left() == area.left() + 12)
            report["checks"]["desktop_" + alignment + "_alignment"] = (
                correct_edge and store.get_setting("desktop_layout", "bottom") == alignment and
                window.settings_alignment_combo.currentData() == alignment)
            first_note.grab().save(str(report_path.with_name(report_path.stem + "-" + alignment + "-note.png")))
        report["checks"]["alignment_leaves_input_notepad_position"] = window.capture_notepad.pos() == original_capture_position
        window.alignment_combo.setCurrentIndex(window.alignment_combo.findData("bottom"))
        window.align_bottom_button.click()
        area = QApplication.primaryScreen().availableGeometry()
        report["checks"]["sticky_windows_inside_desktop"] = all(area.contains(n.geometry()) for n in window.notes.values())
        report["checks"]["bottom_alignment_saved"] = store.get_setting("desktop_layout") == "bottom"
        if window.notes:
            note = next(iter(window.notes.values()))
            if sys.platform == "win32" and app.platformName() == "windows":
                import ctypes
                from ctypes import wintypes
                probe = QWidget()
                probe.setWindowTitle("알리미 실행 검증 기준 창")
                probe.resize(240, 120)
                probe.show()
                probe.raise_()
                note.activateWindow()
                app.processEvents()
                note.keep_below_windows()
                user = ctypes.windll.user32
                user.GetWindowLongPtrW.argtypes = [wintypes.HWND, ctypes.c_int]
                user.GetWindowLongPtrW.restype = ctypes.c_ssize_t
                user.IsWindowVisible.argtypes = [wintypes.HWND]
                user.IsWindowVisible.restype = wintypes.BOOL
                user.GetClassNameW.argtypes = [wintypes.HWND, wintypes.LPWSTR, ctypes.c_int]
                positions = []
                shell_positions = []
                callback_type = ctypes.WINFUNCTYPE(wintypes.BOOL, wintypes.HWND, wintypes.LPARAM)
                @callback_type
                def collect_position(hwnd, _):
                    if user.IsWindowVisible(hwnd):
                        name = ctypes.create_unicode_buffer(256)
                        user.GetClassNameW(hwnd, name, len(name))
                        if name.value in ("Progman", "WorkerW"):
                            shell_positions.append(len(positions))
                        positions.append(int(hwnd))
                    return True
                user.EnumWindows.argtypes = [callback_type, wintypes.LPARAM]
                user.EnumWindows(collect_position, 0)
                note_id, probe_id = int(note.winId()), int(probe.winId())
                report["checks"]["native_note_not_topmost"] = not bool(user.GetWindowLongPtrW(note_id, -20) & 0x8)
                report["checks"]["native_note_below_ordinary_window"] = note_id in positions and probe_id in positions and positions.index(note_id) > positions.index(probe_id)
                report["checks"]["native_note_above_wallpaper"] = note_id in positions and (not shell_positions or positions.index(note_id) < min(shell_positions))
                note.title_label.click()
                app.processEvents()
                window.sync_notes()
                app.processEvents()
                probe.raise_()
                app.processEvents()
                positions.clear()
                shell_positions.clear()
                user.EnumWindows(collect_position, 0)
                report["checks"]["expanded_sticky_is_native_topmost_after_refresh"] = (
                    note.expanded and bool(user.GetWindowLongPtrW(note_id, -20) & 0x8))
                report["checks"]["expanded_sticky_above_ordinary_window"] = (
                    note_id in positions and probe_id in positions and
                    positions.index(note_id) < positions.index(probe_id))
                report["checks"]["expanded_sticky_above_other_notes"] = all(
                    int(other.winId()) not in positions or
                    positions.index(note_id) < positions.index(int(other.winId()))
                    for other in window.notes.values() if other is not note)
                note.grab().save(str(report_path.with_name(report_path.stem + "-foreground-detail.png")))
                note.title_label.click()
                app.processEvents()
                positions.clear()
                shell_positions.clear()
                user.EnumWindows(collect_position, 0)
                report["checks"]["collapsed_sticky_returns_below_ordinary_window"] = (
                    not note.expanded and not bool(user.GetWindowLongPtrW(note_id, -20) & 0x8) and
                    note_id in positions and probe_id in positions and
                    positions.index(note_id) > positions.index(probe_id))
                for collapsed, label in ((True, "icon"), (False, "editor")):
                    capture.set_collapsed(collapsed)
                    window.sync_notes()
                    probe.raise_()
                    app.processEvents()
                    positions.clear()
                    user.EnumWindows(collect_position, 0)
                    capture_id = int(capture.winId())
                    report["checks"][f"native_notepad_{label}_is_topmost"] = bool(
                        user.GetWindowLongPtrW(capture_id, -20) & 0x8)
                    report["checks"][f"native_notepad_{label}_above_ordinary_window"] = (
                        capture_id in positions and probe_id in positions and
                        positions.index(capture_id) < positions.index(probe_id))
                probe.close()
        event_count = len(store.list_events())
        edited_date = (date.today() + timedelta(days=1)).isoformat()
        def save_existing_dialog():
            dialog = app.activeModalWidget()
            dialog.title_field.setText("실행 검증 일정 수정")
            dialog.date_field.setText(edited_date)
            dialog.notes_field.setPlainText("수정 저장 실행 검증")
            dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Save).click()
        QTimer.singleShot(0, save_existing_dialog)
        window.edit_event(event.id)
        app.processEvents()
        saved_edit = store.get_event(event.id)
        report["checks"]["existing_manager_edit_saved_and_selected"] = (
            len(store.list_events()) == event_count and saved_edit.title == "실행 검증 일정 수정" and
            saved_edit.start_date == edited_date and saved_edit.notes == "수정 저장 실행 검증" and
            window.calendar.selectedDate().toString("yyyy-MM-dd") == edited_date and
            window.agenda.currentItem().data(Qt.UserRole) == event.id and
            window.notes[event.id].title_label.full_title == saved_edit.title)
        window.process_import(original_deadline, "기존 일정 수정 실행 검증")
        dialog = window.import_dialog
        dialog.notes_editor.setPlainText("재등록에서 수정한 세부 내용")
        dialog.table.item(0, dialog.END_DATE).setText(f"{date.today().year}-10-31")
        dialog.register_button.click()
        app.processEvents()
        saved_import_edit = store.get_event(deadline_event.id)
        report["checks"]["existing_review_edit_updates_same_event_and_sticky"] = (
            len(store.list_events()) == event_count and
            saved_import_edit.notes == "재등록에서 수정한 세부 내용" and
            saved_import_edit.end_date == f"{date.today().year}-10-31" and
            saved_import_edit.notes in window.notes[deadline_event.id].details_label.text() and
            "수정 1건" in window.import_result.toPlainText())
        history_edit = store.upsert_event(Event(title="완료 기록 순서 변경 검증", start_date="2099-10-14"))
        history_other = store.upsert_event(Event(title="완료 상태 유지 검증", start_date="2099-10-15"))
        store.complete_event(history_edit.id)
        store.complete_event(history_other.id)
        window.refresh()
        history_row = next(row for row in range(window.history_table.rowCount())
                           if window.history_table.item(row, 0).data(Qt.UserRole) == history_edit.id)
        window.history_table.setCurrentCell(history_row, 0)
        def save_history_dialog():
            dialog = app.activeModalWidget()
            dialog.date_field.setText("2099-10-16")
            dialog.findChild(QDialogButtonBox).button(QDialogButtonBox.Save).click()
        QTimer.singleShot(0, save_history_dialog)
        window.edit_event(history_edit.id)
        report["checks"]["history_date_edit_preserves_selected_identity"] = (
            window.history_table.item(window.history_table.currentRow(), 0).data(Qt.UserRole) == history_edit.id and
            store.get_event(history_edit.id).start_date == "2099-10-16")
        window.table_action(window.history_table, "done")
        report["checks"]["complete_cancel_targets_edited_history_event"] = (
            store.get_event(history_edit.id).status == "open" and
            store.get_event(history_other.id).status == "done" and
            window.history_table.currentRow() == -1)
        from contextlib import closing
        from dataclasses import asdict
        def persistent_events(repository):
            return [{key: value for key, value in asdict(event).items() if key not in ("x", "y")}
                    for event in repository.list_events()]
        reset_day = store.upsert_event(Event(title="기간 초기화 날짜 검증", start_date="2098-06-15"))
        reset_month = store.upsert_event(Event(title="기간 초기화 월 검증", start_date="2098-06-01", date_precision="month"))
        store.complete_event(reset_month.id)
        reset_outside = store.upsert_event(Event(title="기간 초기화 밖 검증", start_date="2098-07-01"))
        reset_undated = store.upsert_event(Event(title="전체 초기화 날짜 미정 검증", status="review"))
        window.refresh()
        window.reset_from_date.setDate(QDate(2098, 6, 15))
        window.reset_to_date.setDate(QDate(2098, 6, 15))
        before_reset = persistent_events(store)
        before_backups = set(store.path.parent.glob("backups/*.sqlite3"))
        close_with_answer(window, QMessageBox.Cancel, window.reset_range_button)
        report["checks"]["range_reset_cancel_keeps_all_schedules"] = (
            persistent_events(store) == before_reset and set(store.path.parent.glob("backups/*.sqlite3")) == before_backups)
        close_with_answer(window, QMessageBox.Ok, window.reset_range_button)
        report["checks"]["range_reset_deletes_dated_and_completed_month_only"] = (
            store.get_event(reset_day.id) is None and store.get_event(reset_month.id) is None and
            store.get_event(reset_outside.id) is not None and store.get_event(reset_undated.id) is not None)
        new_backups = set(store.path.parent.glob("backups/*.sqlite3")) - before_backups
        range_backup = next(iter(new_backups))
        with closing(Store(range_backup)) as saved:
            report["checks"]["range_reset_backs_up_original_schedules"] = persistent_events(saved) == before_reset
        all_snapshot = persistent_events(store)
        saved_draft = "일정 전체 초기화에도 유지할 메모 초안"
        window.capture_notepad.editor.setPlainText(saved_draft)
        saved_unlimited = store.get_setting("reminder_unlimited")
        close_with_answer(window, QMessageBox.Ok, window.reset_all_button)
        report["checks"]["all_reset_clears_schedule_views_and_keeps_memo_icon"] = (
            not store.list_events() and not window.notes and
            (window.desktop_list is None or not window.desktop_list.isVisible()) and
            all(window.agenda.item(row).data(Qt.UserRole) is None for row in range(window.agenda.count())) and
            window.review_table.rowCount() == 0 and
            window.history_table.rowCount() == 0 and window.capture_notepad.isVisible())
        report["checks"]["all_reset_preserves_draft_and_display_settings"] = (
            store.get_setting("capture_draft") == saved_draft and
            store.get_setting("reminder_unlimited") == saved_unlimited)
        full_backup = next(iter(set(store.path.parent.glob("backups/*.sqlite3")) - before_backups - new_backups))
        with closing(Store(full_backup)) as saved:
            report["checks"]["all_reset_backs_up_remaining_events"] = persistent_events(saved) == all_snapshot
        store.restore(full_backup)
        window.refresh()
        report["checks"]["reset_backup_restores_events_and_desktop_reminders"] = (
            persistent_events(store) == all_snapshot and
            set(window.notes) == {event.id for event in store.reminder_events()})
        before_exit = persistent_events(store)
        unfinished_before_exit = any(event["status"] == "open" for event in before_exit)
        exit_draft = "트레이 종료 직전 입력한 로컬 보존 검증 초안"
        window.capture_notepad.editor.setPlainText(exit_draft)
        exit_action = next(action for action in window.tray.contextMenu().actions()
                           if action.text() == "알리미 종료")
        exit_timeout = {"fired": False}
        watchdog = QTimer()
        watchdog.setSingleShot(True)
        def stop_failed_exit():
            exit_timeout["fired"] = True
            app.quit()
        watchdog.timeout.connect(stop_failed_exit)
        watchdog.start(5000)
        QTimer.singleShot(0, exit_action.trigger)
        app.exec_()
        watchdog.stop()
        report["checks"]["tray_exit_accepts_unfinished_events"] = (
            unfinished_before_exit and app.property("shutting_down") and not exit_timeout["fired"])
        reloaded = Store(store.path)
        try:
            report["checks"]["tray_exit_preserves_local_events_and_status"] = persistent_events(reloaded) == before_exit
            report["checks"]["tray_exit_flushes_recent_local_draft"] = reloaded.get_setting("capture_draft") == exit_draft
        finally:
            reloaded.close()
        report["success"] = all(report["checks"].values())
    except Exception:
        report["error"] = traceback.format_exc()
    finally:
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        app.setProperty("shutting_down", True)
        for note in window.notes.values():
            note.allow_close = True
            note.close()
        if window.desktop_list:
            window.desktop_list.allow_close = True
            window.desktop_list.close()
        window.on_shutdown()
    return 0 if report["success"] else 1


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--background", action="store_true")
    parser.add_argument("--no-autostart", action="store_true")
    parser.add_argument("--data-dir", type=Path)
    parser.add_argument("--smoke-test", type=Path)
    parser.add_argument("--smoke-document", type=Path)
    args = parser.parse_args()
    frozen_install()
    args.data_dir = args.data_dir or default_data_dir()
    args.data_dir.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(filename=str(args.data_dir / "app.log"), level=logging.WARNING,
                        format="%(asctime)s %(levelname)s %(message)s", encoding="utf-8")
    QApplication.setAttribute(Qt.AA_EnableHighDpiScaling, True)
    QApplication.setAttribute(Qt.AA_UseHighDpiPixmaps, True)
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setOrganizationName("SchoolAlimi")
    app.setQuitOnLastWindowClosed(False)
    app.setStyleSheet(STYLE)
    icon = make_icon()
    app.setWindowIcon(icon)
    server_name = "SchoolAlimi-" + hashlib.sha256(str(args.data_dir.resolve()).encode()).hexdigest()[:16]
    socket = QLocalSocket()
    socket.connectToServer(server_name)
    if socket.waitForConnected(400):
        socket.write(b"version\n")
        socket.waitForBytesWritten(500)
        if socket.waitForReadyRead(1500):
            existing_version = bytes(socket.readAll()).decode("utf-8", errors="replace").strip()
        else:
            existing_version = ""
        socket.disconnectFromServer()
        if existing_version != APP_VERSION:
            QMessageBox.warning(None, APP_NAME,
                "기존 버전이 아직 실행 중이어서 새 버전을 열지 못했습니다.\n"
                "자료 분석이 끝난 뒤 기존 알리미를 종료하고 다시 실행해 주세요.\n\n"
                f"실행하려는 버전: {APP_VERSION} · 기존 버전: {existing_version or '0.1.0~0.1.3'}")
            return 1
        show_socket = QLocalSocket()
        show_socket.connectToServer(server_name)
        if show_socket.waitForConnected(400):
            show_socket.write(b"show\n")
            show_socket.waitForBytesWritten(500)
        return 0
    QLocalServer.removeServer(server_name)
    server = QLocalServer()
    if not server.listen(server_name):
        QMessageBox.warning(None, APP_NAME, "알리미를 실행할 수 없습니다. 기존 실행 여부를 확인해 주세요.")
        return 1
    try:
        store = Store(args.data_dir / "alimi.sqlite3")
        window = MainWindow(store, args, icon)
        def incoming():
            connection = server.nextPendingConnection()
            if connection:
                handled = False
                command_buffer = bytearray()
                def receive_command():
                    nonlocal handled
                    if handled:
                        return
                    command_buffer.extend(bytes(connection.readAll()))
                    if b"\n" not in command_buffer and command_buffer != b"show":
                        return
                    command = bytes(command_buffer).strip()
                    handled = True
                    if command == b"version":
                        connection.write((APP_VERSION + "\n").encode("utf-8"))
                        connection.flush()
                    else:
                        window.show_manager()
                    connection.disconnectFromServer()
                connection.readyRead.connect(receive_command)
                connection.disconnected.connect(connection.deleteLater)
                receive_command()
        server.newConnection.connect(incoming)
        if args.smoke_test:
            return smoke_test(app, window, store, args.smoke_test)
        if not args.background:
            window.show()
        result = app.exec_()
        server.close()
        return result
    except Exception as exc:
        logging.exception("Startup failed")
        if args.smoke_test:
            args.smoke_test.write_text(json.dumps({"success": False, "error": traceback.format_exc()}, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            QMessageBox.critical(None, APP_NAME, "앱을 시작하지 못했습니다.\n" + str(exc))
        return 1
