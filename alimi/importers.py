"""Read school documents locally, without executing embedded content.

The optional OCR callback receives a Pillow image and returns recognized text.
All sources are opened read-only; extraction never changes the original file.
"""

from __future__ import annotations

import datetime as dt
import re
import struct
import warnings
import xml.etree.ElementTree as ET
import zipfile
import zlib
from pathlib import Path
from typing import Callable

from PIL import Image


class DocumentImportError(ValueError):
    """An extraction failure suitable for displaying to the user."""


SUPPORTED_EXTENSIONS = frozenset(
    {".txt", ".csv", ".tsv", ".md", ".docx", ".xlsx", ".pdf", ".hwpx", ".hwp",
     ".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
)
MAX_FILE_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_BYTES = 200 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 5000
MAX_TEXT_CHARS = 4_000_000
MAX_DOCUMENT_PAGES = 300
MAX_IMAGE_PIXELS = 40_000_000
OCRFunction = Callable[[Image.Image], str]


def extract_file_text(path: str | Path, ocr_func: OCRFunction | None = None) -> str:
    """Extract text from a supported local file, optionally using local OCR.

    DocumentImportError has a Korean message for unsupported, encrypted, empty,
    damaged, or excessively large files. No shell, Office, or network is used.
    """
    source = Path(path)
    try:
        if not source.is_file():
            raise DocumentImportError("파일을 찾을 수 없습니다. 파일 위치를 확인해 주세요.")
        if source.stat().st_size > MAX_FILE_BYTES:
            raise DocumentImportError("파일이 너무 큽니다. 100MB 이하의 파일을 가져와 주세요.")
        suffix = source.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise DocumentImportError(
                "지원하지 않는 파일 형식입니다. TXT, DOCX, XLSX, PDF, HWPX, HWP 또는 이미지를 선택해 주세요."
            )
        if suffix in {".docx", ".xlsx", ".hwpx"}:
            _validate_archive(source)
        if suffix in {".txt", ".csv", ".tsv", ".md"}:
            text = _read_plain_text(source)
        elif suffix == ".docx":
            text = _read_docx(source)
        elif suffix == ".xlsx":
            text = _read_xlsx(source)
        elif suffix == ".pdf":
            text = _read_pdf(source, ocr_func)
        elif suffix == ".hwpx":
            text = _read_hwpx(source)
        elif suffix == ".hwp":
            text = _read_hwp(source)
        else:
            text = _read_image(source, ocr_func)
        text = _clean_text(text)
        if not text:
            raise DocumentImportError(
                "읽을 수 있는 글자를 찾지 못했습니다. 스캔 자료라면 선명한 이미지나 텍스트로 다시 가져와 주세요."
            )
        return text
    except DocumentImportError:
        raise
    except PermissionError as exc:
        raise DocumentImportError("파일을 읽을 수 없습니다. 접근 권한이나 다른 프로그램의 사용 여부를 확인해 주세요.") from exc
    except ImportError as exc:
        raise DocumentImportError("이 파일을 읽는 구성 요소가 없습니다. 알리미를 다시 설치해 주세요.") from exc
    except Exception as exc:
        raise DocumentImportError(
            "파일을 읽지 못했습니다. 손상되거나 암호로 보호된 파일인지 확인하고, 다른 형식으로 저장해 주세요."
        ) from exc


def _clean_text(text: str) -> str:
    if not isinstance(text, str):
        raise DocumentImportError("문자 인식 결과가 올바르지 않습니다. 다시 가져와 주세요.")
    if len(text) > MAX_TEXT_CHARS:
        raise DocumentImportError("문서의 글자 수가 너무 많습니다. 필요한 부분만 나누어 가져와 주세요.")
    return text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n").strip()


def _validate_archive(path: Path) -> None:
    """Validate containers before document libraries inflate their contents."""
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > MAX_ARCHIVE_MEMBERS:
            raise DocumentImportError("문서의 내부 파일 수가 너무 많습니다. 필요한 부분만 가져와 주세요.")
        total = 0
        for member in members:
            if member.flag_bits & 1:
                raise DocumentImportError("암호로 보호된 문서입니다. 암호를 해제한 사본을 가져와 주세요.")
            total += member.file_size
            # No files are extracted to disk. Reject excessive expansion anyway.
            if total > MAX_ARCHIVE_BYTES or member.file_size > MAX_FILE_BYTES:
                raise DocumentImportError("압축을 해제한 문서가 너무 큽니다. 필요한 부분만 가져와 주세요.")
            if member.file_size > 1024 * 1024 and member.file_size / max(1, member.compress_size) > 1000:
                raise DocumentImportError("문서 압축 비율이 비정상적으로 높아 읽기를 중단했습니다.")


def _read_plain_text(path: Path) -> str:
    data = path.read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    for encoding in ("utf-8-sig", "cp949"):
        try:
            result = data.decode(encoding)
            if "\x00" in result:
                raise DocumentImportError("텍스트 파일에 읽을 수 없는 데이터가 포함되어 있습니다.")
            return result
        except UnicodeDecodeError:
            continue
    raise DocumentImportError("텍스트 인코딩을 읽지 못했습니다. UTF-8 또는 CP949로 저장해 주세요.")


def _read_docx(path: Path) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    lines: list[str] = []
    for child in document.element.body:
        kind = _local_name(child.tag)
        if kind == "p":
            lines.append(Paragraph(child, document).text)
        elif kind == "tbl":
            table = Table(child, document)
            lines.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
    return "\n".join(lines)


def _spreadsheet_value(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0):
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.strftime("%H:%M:%S")
    return str(value)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True, keep_links=False)
    lines: list[str] = []
    total = 0
    try:
        for sheet in workbook.worksheets:
            if (sheet.max_row or 0) > 100_000 or (sheet.max_column or 0) > 2000:
                raise DocumentImportError("스프레드시트의 행이나 열이 너무 많습니다. 필요한 범위만 가져와 주세요.")
            lines.append(f"[{sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                values = [_spreadsheet_value(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    line = "\t".join(values)
                    total += len(line)
                    if total > MAX_TEXT_CHARS:
                        raise DocumentImportError("스프레드시트가 너무 큽니다. 필요한 시트나 행만 가져와 주세요.")
                    lines.append(line)
    finally:
        workbook.close()
    return "\n".join(lines)


def _recognize(image: Image.Image, ocr_func: OCRFunction | None) -> str:
    if ocr_func is None:
        raise DocumentImportError("이미지를 읽으려면 로컬 문자 인식 기능이 필요합니다. 텍스트를 직접 붙여넣을 수도 있습니다.")
    if image.width * image.height > MAX_IMAGE_PIXELS:
        raise DocumentImportError("이미지 해상도가 너무 큽니다. 필요한 영역만 캡처해 주세요.")
    try:
        return _clean_text(ocr_func(image))
    except DocumentImportError:
        raise
    except Exception as exc:
        raise DocumentImportError("로컬 문자 인식에 실패했습니다. 선명한 이미지로 다시 시도해 주세요.") from exc


def _read_image(path: Path, ocr_func: OCRFunction | None) -> str:
    # Treat Pillow's decompression bomb warning as an error, never globally
    # disable its protections. TIFF may contain multiple scanned pages.
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(path) as source:
            page_count = getattr(source, "n_frames", 1)
            if page_count > MAX_DOCUMENT_PAGES:
                raise DocumentImportError("이미지의 페이지 수가 너무 많습니다. 300페이지 이하로 나누어 가져와 주세요.")
            lines: list[str] = []
            for index in range(page_count):
                source.seek(index)
                if source.width * source.height > MAX_IMAGE_PIXELS:
                    raise DocumentImportError("이미지 해상도가 너무 큽니다. 필요한 영역만 캡처해 주세요.")
                with source.convert("RGB") as page_image:
                    lines.append(_recognize(page_image, ocr_func))
            return "\n\n".join(lines)


def _read_pdf(path: Path, ocr_func: OCRFunction | None) -> str:
    import pdfplumber

    page_texts: list[str] = []
    scanned_pages: list[int] = []
    table_headers: set[tuple[str, ...]] = set()
    with pdfplumber.open(str(path)) as document:
        if len(document.pages) > MAX_DOCUMENT_PAGES:
            raise DocumentImportError("PDF가 너무 깁니다. 300페이지 이하로 나누어 가져와 주세요.")
        for index, page in enumerate(document.pages):
            content = page.extract_text() or ""
            # A scan may have a tiny selectable page number or header. OCR that
            # page when available, while retaining text on genuine text pages.
            if not content.strip() or (page.images and len(re.sub(r"\s", "", content)) < 20):
                scanned_pages.append(index)
                page_texts.append(content)
            else:
                page_texts.append(_pdf_page_text(page, content, table_headers))
    if scanned_pages:
        if ocr_func is None:
            raise DocumentImportError("PDF에 스캔된 페이지가 있습니다. 로컬 문자 인식 기능을 준비하거나 텍스트를 붙여넣어 주세요.")
        import pypdfium2 as pdfium

        document = pdfium.PdfDocument(str(path))
        try:
            for index in scanned_pages:
                page = document[index]
                bitmap = None
                try:
                    width, height = page.get_size()
                    scale = min(2.5, (MAX_IMAGE_PIXELS * 0.95 / max(1, width * height)) ** 0.5)
                    bitmap = page.render(scale=scale)
                    with bitmap.to_pil().convert("RGB") as page_image:
                        recognized = _recognize(page_image, ocr_func)
                    page_texts[index] = recognized or page_texts[index]
                finally:
                    if bitmap is not None:
                        bitmap.close()
                    page.close()
        finally:
            document.close()
    return "\n\n".join(page_texts)


def _pdf_table_rows(table) -> list[list[str]]:
    """Keep each ruled-table row together and restore actual vertical merges.

    pdfplumber represents a merged cell in later rows as None. A blank cell is
    different: it has its own rectangle and must not inherit a previous value.
    Reuse a value only when the earlier cell's rectangle covers this row.
    """
    extracted = table.extract()
    merged: dict[int, tuple[tuple[float, float, float, float], str]] = {}
    rows: list[list[str]] = []
    for values, row in zip(extracted, table.rows):
        rectangles = [cell for cell in row.cells if cell is not None]
        if not rectangles:
            continue
        row_top = min(cell[1] for cell in rectangles)
        row_bottom = min(cell[3] for cell in rectangles if cell[3] > row_top)
        cells: list[str] = []
        for column, value in enumerate(values):
            rectangle = row.cells[column]
            text = re.sub(r"\s+", " ", value or "").strip()
            if rectangle is not None:
                merged[column] = (rectangle, text)
            elif column in merged:
                previous_rectangle, previous_text = merged[column]
                # A missing cell from a horizontal merge or table detection
                # error must not copy an unrelated value into a new row.
                if (previous_rectangle[1] < row_top
                        and previous_rectangle[3] >= row_bottom - 0.01):
                    text = previous_text
            cells.append(text)
        if any(cells):
            rows.append(cells)
    return rows


def _pdf_schedule_header(row: list[str]) -> bool:
    normalized = {re.sub(r"\s", "", value) for value in row}
    time_labels = {"시기", "평가시기", "평가월", "날짜", "일자", "기간", "월"}
    content_labels = {"교과", "단원명", "평가내용", "제목", "내용", "행사", "업무", "구분"}
    return bool(normalized & time_labels and normalized & content_labels)


def _pdf_page_text(page, plain_text: str, seen_headers: set[tuple[str, ...]]) -> str:
    """Prefer selectable text while preserving columns in recognized tables.

    Table characters are removed from ordinary page text before structured
    rows are inserted, so a single assessment cannot become duplicate events.
    Text above/below a table (especially its academic year) keeps its order.
    Unruled or unrecognized tables keep the existing plain-text extraction.
    """
    tables = []
    for table in page.find_tables():
        rows = _pdf_table_rows(table)
        # Very small rectangles in a title/logo are not useful tabular data.
        if len(rows) >= 2 and max(len(row) for row in rows) >= 2:
            tables.append((table, rows))
    if not tables:
        return plain_text

    def outside_table(obj):
        if obj.get("object_type") != "char":
            return True
        x = (obj["x0"] + obj["x1"]) / 2
        y = (obj["top"] + obj["bottom"]) / 2
        return not any(table.bbox[0] <= x <= table.bbox[2]
                       and table.bbox[1] <= y <= table.bbox[3]
                       for table, _ in tables)

    blocks = [(line["top"], line["x0"], line["text"])
              for line in page.filter(outside_table).extract_text_lines(return_chars=False)]
    for table, rows in tables:
        lines = []
        for row in rows:
            if _pdf_schedule_header(row):
                key = tuple(re.sub(r"\s", "", value) for value in row)
                if key in seen_headers:
                    continue
                seen_headers.add(key)
            lines.append(" | ".join(row))
        if lines:
            blocks.append((table.bbox[1], table.bbox[0], "[표]\n" + "\n".join(lines)))
    return "\n".join(text for _, _, text in sorted(blocks, key=lambda block: block[:2]))


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _xml_text(node: ET.Element) -> str:
    """Inline paragraph content, excluding nested table paragraphs."""
    kind = _local_name(node.tag)
    if kind == "tbl":
        return ""
    if kind == "lineBreak":
        return "\n"
    if kind == "tab":
        return "\t"
    if kind == "t":
        return "".join(node.itertext())
    return "".join(_xml_text(child) for child in node)


def _hwpx_blocks(node: ET.Element) -> list[str]:
    kind = _local_name(node.tag)
    if kind == "tbl":
        rows = []
        for row in node:
            if _local_name(row.tag) != "tr":
                continue
            cells = ["\n".join(_hwpx_blocks(cell)) for cell in row if _local_name(cell.tag) == "tc"]
            rows.append("\t".join(cells))
        return rows
    if kind == "p":
        lines = [_xml_text(node)]
        # A HWPX table is normally inside a run in an otherwise empty p.
        def nested_tables(element: ET.Element):
            for child in element:
                if _local_name(child.tag) == "tbl":
                    yield child
                else:
                    yield from nested_tables(child)
        for table in nested_tables(node):
            lines.extend(_hwpx_blocks(table))
        return [line for line in lines if line]
    return [line for child in node for line in _hwpx_blocks(child)]


def _read_hwpx(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        section_names = [name for name in archive.namelist() if re.fullmatch(r"Contents/section\d+\.xml", name, re.I)]
        section_names.sort(key=lambda name: int(re.search(r"section(\d+)\.xml", name, re.I).group(1)))
        if not section_names:
            raise DocumentImportError("HWPX 본문을 찾지 못했습니다. 한글에서 HWPX로 다시 저장해 주세요.")
        lines: list[str] = []
        for name in section_names:
            data = archive.read(name)
            upper = data.upper()
            if b"<!DOCTYPE" in upper or b"<!ENTITY" in upper:
                raise DocumentImportError("외부 참조가 있는 XML 문서는 읽을 수 없습니다. 일반 HWPX로 다시 저장해 주세요.")
            root = ET.fromstring(data)
            lines.extend(_hwpx_blocks(root))
        return "\n".join(lines)


def _inflate_hwp(data: bytes) -> bytes:
    inflater = zlib.decompressobj(-15)
    decoded = inflater.decompress(data, MAX_FILE_BYTES + 1)
    if len(decoded) > MAX_FILE_BYTES or inflater.unconsumed_tail:
        raise DocumentImportError("HWP 본문의 압축 해제 크기가 너무 큽니다.")
    decoded += inflater.flush()
    if len(decoded) > MAX_FILE_BYTES or not inflater.eof:
        raise DocumentImportError("HWP 본문이 손상되었거나 너무 큽니다.")
    return decoded


def _hwp_paragraph(data: bytes) -> str:
    encoded = bytearray()
    i = 0
    unit_count = len(data) // 2
    extended_controls = set(range(1, 10)) | {11, 12} | set(range(14, 24))
    while i < unit_count:
        char = struct.unpack_from("<H", data, i * 2)[0]
        if char in extended_controls:
            if char == 9:
                encoded.extend(b"\t\x00")
            i += 8
            continue
        if char in (10, 13):
            encoded.extend(b"\n\x00")
        elif char in (30, 31):
            encoded.extend(b" \x00")
        elif char >= 32:
            encoded.extend(struct.pack("<H", char))
        if len(encoded) > MAX_TEXT_CHARS * 2:
            raise DocumentImportError("HWP 문단의 글자 수가 너무 많습니다. 필요한 부분만 가져와 주세요.")
        i += 1
    return encoded.decode("utf-16-le", errors="replace").strip()


def _hwp_records(data: bytes) -> list[str]:
    offset = 0
    lines: list[str] = []
    while offset + 4 <= len(data):
        header = struct.unpack_from("<I", data, offset)[0]
        offset += 4
        tag = header & 0x3FF
        size = header >> 20
        if size == 0xFFF:
            if offset + 4 > len(data):
                raise DocumentImportError("HWP 본문이 손상되었습니다. HWPX로 다시 저장해 주세요.")
            size = struct.unpack_from("<I", data, offset)[0]
            offset += 4
        if offset + size > len(data):
            raise DocumentImportError("HWP 본문이 손상되었습니다. HWPX로 다시 저장해 주세요.")
        if tag == 67:  # HWPTAG_PARA_TEXT
            text = _hwp_paragraph(data[offset:offset + size])
            if text:
                lines.append(text)
        offset += size
    return lines


def _read_hwp(path: Path) -> str:
    import olefile

    with olefile.OleFileIO(str(path)) as document:
        if not document.exists("FileHeader"):
            raise DocumentImportError("지원하는 HWP 5 문서가 아닙니다. HWPX로 저장해 주세요.")
        header = document.openstream("FileHeader").read(256)
        if len(header) < 40 or not header.startswith(b"HWP Document File"):
            raise DocumentImportError("지원하는 HWP 5 문서가 아닙니다. HWPX로 저장해 주세요.")
        flags = struct.unpack_from("<I", header, 36)[0]
        if flags & 2:
            raise DocumentImportError("암호로 보호된 HWP입니다. 암호를 해제한 사본을 가져와 주세요.")
        if flags & 4:
            raise DocumentImportError("배포용 HWP입니다. 편집 가능한 HWPX 또는 PDF로 저장해 주세요.")
        sections = [entry for entry in document.listdir() if len(entry) == 2 and entry[0] == "BodyText" and re.fullmatch(r"Section\d+", entry[1])]
        sections.sort(key=lambda entry: int(entry[1][7:]))
        lines: list[str] = []
        total_size = 0
        for section in sections:
            data = document.openstream(section).read(MAX_FILE_BYTES + 1)
            if len(data) > MAX_FILE_BYTES:
                raise DocumentImportError("HWP 본문이 너무 큽니다. 필요한 부분만 가져와 주세요.")
            if flags & 1:
                data = _inflate_hwp(data)
            total_size += len(data)
            if total_size > MAX_ARCHIVE_BYTES:
                raise DocumentImportError("HWP 본문이 너무 큽니다. 필요한 부분만 가져와 주세요.")
            lines.extend(_hwp_records(data))
        if not lines and document.exists("PrvText"):
            preview = document.openstream("PrvText").read(MAX_FILE_BYTES + 1)
            return preview.decode("utf-16-le", errors="replace")
        return "\n".join(lines)
