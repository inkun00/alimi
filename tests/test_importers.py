"""Integration checks with generated, disposable local document fixtures."""

import datetime as dt
import struct
import tempfile
import unittest
import zipfile
import zlib
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image
from docx import Document
from openpyxl import Workbook

from alimi.importers import (
    DocumentImportError, _hwp_records, _inflate_hwp, _pdf_page_text,
    _pdf_table_rows, extract_file_text,
)


def write_pdf(path: Path, text: str | None) -> None:
    """Small valid fixture PDF, requiring no proprietary document tooling."""
    content = b"" if text is None else f"BT /F1 12 Tf 40 750 Td ({text}) Tj ET".encode("ascii")
    write_pdf_content(path, content)


def write_pdf_content(path: Path, content: bytes) -> None:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    path.write_bytes(data)


def fake_table(values, rectangles=None, top=20):
    """Coordinate-aware fixture without any personal school document data."""
    if rectangles is None:
        rectangles = [[(column * 100, top + index * 20,
                        (column + 1) * 100, top + (index + 1) * 20)
                       for column in range(len(row))]
                      for index, row in enumerate(values)]
    return SimpleNamespace(
        extract=lambda: values,
        rows=[SimpleNamespace(cells=row) for row in rectangles],
        bbox=(0, top, len(values[0]) * 100, top + len(values) * 20),
    )


class ImporterTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent)
        self.root = Path(self.temp.name)
        self.previous_tempdir = tempfile.tempdir
        tempfile.tempdir = str(self.root)

    def tearDown(self):
        tempfile.tempdir = self.previous_tempdir
        self.temp.cleanup()

    def test_text_encodings_and_lines(self):
        for encoding in ("utf-8-sig", "cp949", "utf-16"):
            with self.subTest(encoding=encoding):
                path = self.root / "school.txt"
                path.write_bytes("교직원 회의\r\n2026년 10월 13일 15:30".encode(encoding))
                self.assertEqual(extract_file_text(path), "교직원 회의\n2026년 10월 13일 15:30")

    def test_docx_keeps_paragraph_table_order(self):
        document = Document()
        document.add_paragraph("학교 공지")
        row = document.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = "2026년 10월 13일"
        row.cells[1].text = "교직원 회의"
        document.add_paragraph("장소: 2층 회의실")
        path = self.root / "school.docx"
        document.save(path)
        self.assertEqual(extract_file_text(path), "학교 공지\n2026년 10월 13일\t교직원 회의\n장소: 2층 회의실")

    def test_xlsx_date_time_and_formula_never_executed(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "업무일정"
        sheet.append(["제목", "날짜", "시간"])
        sheet.append(["교직원 회의", dt.datetime(2026, 10, 13), dt.time(15, 30)])
        sheet.append(["=HYPERLINK(\"https://invalid.local\",\"test\")"])
        path = self.root / "school.xlsx"
        workbook.save(path)
        text = extract_file_text(path)
        self.assertIn("[업무일정]", text)
        self.assertIn("교직원 회의\t2026-10-13\t15:30:00", text)
        self.assertNotIn("invalid.local", text)

    def test_hwpx_paragraph_and_table(self):
        path = self.root / "school.hwpx"
        xml = '''<hs:sec xmlns:hs="urn:section" xmlns:hp="urn:paragraph">
          <hp:p><hp:run><hp:t>학교 일정</hp:t></hp:run></hp:p>
          <hp:p><hp:run><hp:tbl><hp:tr>
            <hp:tc><hp:subList><hp:p><hp:run><hp:t>2026년 10월 13일</hp:t></hp:run></hp:p></hp:subList></hp:tc>
            <hp:tc><hp:subList><hp:p><hp:run><hp:t>교직원 회의</hp:t></hp:run></hp:p></hp:subList></hp:tc>
          </hp:tr></hp:tbl></hp:run></hp:p>
        </hs:sec>'''
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("Contents/section0.xml", xml)
        self.assertEqual(extract_file_text(path), "학교 일정\n2026년 10월 13일\t교직원 회의")

    def test_hwpx_rejects_entity_declarations(self):
        path = self.root / "external.hwpx"
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("Contents/section0.xml", '<!DOCTYPE x [<!ENTITY y SYSTEM "file:///secret">]><x/>')
        with self.assertRaisesRegex(DocumentImportError, "외부 참조"):
            extract_file_text(path)

    def test_pdf_text_first_avoids_ocr(self):
        path = self.root / "school.pdf"
        write_pdf(path, "Meeting 2026-10-13 15:30")
        def forbidden_ocr(image):
            self.fail("Text PDF should not require OCR")
        self.assertEqual(extract_file_text(path, forbidden_ocr), "Meeting 2026-10-13 15:30")

    def test_pdf_ruled_table_keeps_row_month_and_merged_subject(self):
        # Native PDF text and rules exercise pdfplumber's actual geometry,
        # including a vertically merged subject and wrapped unit/content.
        path = self.root / "assessment.pdf"
        commands = ["0.5 w"]
        for y in (730, 700, 620):
            commands.append(f"40 {y} m 540 {y} l S")
        commands.append("120 660 m 540 660 l S")
        for x in (40, 120, 220, 430, 540):
            commands.append(f"{x} 730 m {x} 620 l S")
        def text(value, x, y):
            commands.append(f"BT /F1 10 Tf {x} {y} Td ({value}) Tj ET")
        text("2026 assessment plan", 40, 760)
        for value, x in zip(("Subject", "Unit", "Content", "Month"), (50, 130, 230, 440)):
            text(value, x, 712)
        text("Science", 50, 655)
        text("2. Biology", 130, 683)
        text("and ecology", 130, 670)
        text("Compare habitats", 230, 683)
        text("in our region", 230, 670)
        text("September", 440, 683)
        text("4. Climate", 130, 642)
        text("Research weather", 230, 642)
        text("October", 440, 642)
        text("School notice", 40, 595)
        write_pdf_content(path, "\n".join(commands).encode("ascii"))
        def forbidden_ocr(image):
            self.fail("A selectable table must not invoke OCR")
        extracted = extract_file_text(path, forbidden_ocr)
        self.assertIn("Science | 2. Biology and ecology | Compare habitats in our region | September", extracted)
        self.assertIn("Science | 4. Climate | Research weather | October", extracted)
        self.assertEqual(extracted.count("Compare habitats"), 1)
        self.assertLess(extracted.index("2026 assessment plan"), extracted.index("[표]"))
        self.assertGreater(extracted.index("School notice"), extracted.index("October"))

    def test_pdf_merges_use_geometry_and_blank_cells_stay_blank(self):
        values = [["교과", "내용", "시기"], ["국어", "읽기\n평가", "9월"],
                  [None, "쓰기 평가", "10월"], ["", "발표 평가", "11월"],
                  [None, "노래 평가", "12월"]]
        table = fake_table(values)
        table.rows[1].cells[0] = (0, 40, 100, 80)
        table.rows[2].cells[0] = None
        table.rows[4].cells[0] = None
        rows = _pdf_table_rows(table)
        self.assertEqual(rows[1], ["국어", "읽기 평가", "9월"])
        self.assertEqual(rows[2], ["국어", "쓰기 평가", "10월"])
        self.assertEqual(rows[3][0], "")
        self.assertEqual(rows[4][0], "")

    def test_pdf_removes_repeat_schedule_headers_across_pages(self):
        header = ["교과", "영역", "단원명", "평가내용", "평가\n방법", "평가\n시기"]
        seen = set()
        for subject, month in (("국어", "9월"), ("과학", "10월")):
            table = fake_table([header, [subject, "읽기", "1. 계획", "실천하기", "관찰", month]])
            page = SimpleNamespace(
                find_tables=lambda: [table],
                filter=lambda predicate: SimpleNamespace(
                    extract_text_lines=lambda return_chars=False: [
                        {"top": 0, "x0": 0, "text": "2026학년도 평가 계획"}
                    ]),
            )
            result = _pdf_page_text(page, "plain table", seen)
            if subject == "국어":
                self.assertIn("교과 | 영역 | 단원명 | 평가내용 | 평가 방법 | 평가 시기", result)
            else:
                self.assertNotIn("평가 시기", result)
            self.assertIn(subject + " | 읽기 | 1. 계획 | 실천하기 | 관찰 | " + month, result)

    def test_pdf_scan_renders_to_local_ocr(self):
        path = self.root / "scanned.pdf"
        write_pdf(path, None)
        seen = []
        def local_ocr(image):
            self.assertIsInstance(image, Image.Image)
            seen.append(image.size)
            return "2026년 10월 13일 교직원 회의"
        self.assertIn("교직원 회의", extract_file_text(path, local_ocr))
        self.assertEqual(len(seen), 1)
        with self.assertRaisesRegex(DocumentImportError, "문자 인식"):
            extract_file_text(path)

    def test_image_ocr_receives_pillow_image(self):
        path = self.root / "capture.png"
        Image.new("RGB", (200, 100), "white").save(path)
        def local_ocr(image):
            self.assertEqual(image.mode, "RGB")
            return "공문\n2026-10-13 15:30"
        self.assertEqual(extract_file_text(path, local_ocr), "공문\n2026-10-13 15:30")

    def test_size_limit_and_unsupported_file(self):
        path = self.root / "school.txt"
        path.write_text("schedule", encoding="utf-8")
        with patch("alimi.importers.MAX_FILE_BYTES", 1):
            with self.assertRaisesRegex(DocumentImportError, "너무 큽니다"):
                extract_file_text(path)
        unsupported = self.root / "school.exe"
        unsupported.write_bytes(b"MZ")
        with self.assertRaisesRegex(DocumentImportError, "지원하지 않는"):
            extract_file_text(unsupported)

    def test_archive_expansion_limit(self):
        path = self.root / "school.hwpx"
        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("Contents/section0.xml", "x" * 10000)
        with patch("alimi.importers.MAX_ARCHIVE_BYTES", 100):
            with self.assertRaisesRegex(DocumentImportError, "너무 큽니다"):
                extract_file_text(path)

    def test_hwp_paragraph_controls_and_compression(self):
        # UTF-16 paragraph text record containing an 8-code-unit tab control.
        payload = "회의".encode("utf-16-le") + struct.pack("<8H", 9, 0, 0, 0, 0, 0, 0, 0) + "2026-10-13\r".encode("utf-16-le")
        body = struct.pack("<I", 67 | (len(payload) << 20)) + payload
        compressor = zlib.compressobj(wbits=-15)
        compressed = compressor.compress(body) + compressor.flush()
        self.assertEqual(_hwp_records(_inflate_hwp(compressed)), ["회의\t2026-10-13"])

    def test_empty_or_corrupt_documents_fail_clearly(self):
        path = self.root / "empty.txt"
        path.write_text(" \n", encoding="utf-8")
        with self.assertRaisesRegex(DocumentImportError, "글자를 찾지 못했습니다"):
            extract_file_text(path)
        path = self.root / "broken.pdf"
        path.write_text("not a PDF", encoding="ascii")
        with self.assertRaisesRegex(DocumentImportError, "손상되거나 암호"):
            extract_file_text(path)


if __name__ == "__main__":
    unittest.main()
