"""Document regressions found while auditing other schedule layouts."""

import csv
import datetime as dt
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from alimi.core import parse_schedule_text
from alimi.importers import DocumentImportError, extract_file_text


def write_text_pdf(path, rows):
    commands = []
    for x, y, value in rows:
        commands.append(f"BT /F1 10 Tf {x} {y} Td ({value}) Tj ET")
    content = "\n".join(commands).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{index} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    path.write_bytes(data)


class FormatImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=Path(__file__).resolve().parent)
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_csv_quoted_newline_keeps_title_and_time_on_same_event(self):
        path = self.root / "schedule.csv"
        with path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["날짜", "제목", "시간"])
            writer.writerow(["2026-11-02", "교직원 회의,\n2학기 운영 협의", "15:30"])
            writer.writerow(["2026-11-04", "학부모 상담", "09:00"])
        text = extract_file_text(path)
        events = parse_schedule_text(text)
        self.assertEqual(len(events), 2)
        self.assertEqual((events[0].start_date, events[0].time, events[0].title),
                         ("2026-11-02", "15:30", "교직원 회의, 2학기 운영 협의"))
        self.assertEqual(events[1].time, "09:00")
        self.assertNotIn('"', text)

    def test_tsv_quoted_multiline_field_is_one_row(self):
        path = self.root / "schedule.tsv"
        path.write_text('날짜\t내용\n2026-11-02\t"교직원 회의\n운영 협의"', encoding="utf-16")
        self.assertEqual(extract_file_text(path), "날짜\t내용\n2026-11-02\t교직원 회의 운영 협의")

    def test_malformed_quoted_csv_has_actionable_error(self):
        path = self.root / "schedule.csv"
        path.write_text('2026-11-02,"끝나지 않은 제목', encoding="utf-8")
        with self.assertRaisesRegex(DocumentImportError, "따옴표"):
            extract_file_text(path)

    def test_docx_wrapped_cell_does_not_split_schedule_or_lose_time(self):
        document = Document()
        row = document.add_table(rows=1, cols=3).rows[0]
        for cell, value in zip(row.cells, ["2026-11-02", "교직원 회의\n2학기 운영 협의", "15:30"]):
            cell.text = value
        path = self.root / "schedule.docx"
        document.save(path)
        event, = parse_schedule_text(extract_file_text(path))
        self.assertEqual(event.title, "교직원 회의 2학기 운영 협의")
        self.assertEqual(event.time, "15:30")

    def test_xlsx_wrapped_cell_retains_native_date_and_time(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["날짜", "내용", "시간"])
        sheet.append([dt.date(2026, 11, 2), "교직원 회의\n2학기 운영 협의", dt.time(15, 30)])
        path = self.root / "schedule.xlsx"
        workbook.save(path)
        workbook.close()
        event, = parse_schedule_text(extract_file_text(path))
        self.assertEqual((event.start_date, event.time, event.title),
                         ("2026-11-02", "15:30", "교직원 회의 2학기 운영 협의"))

    def test_hwpx_wrapped_cell_keeps_complete_description(self):
        path = self.root / "schedule.hwpx"
        xml = '''<sec><p><tbl><tr>
          <tc><p><t>2026-11-02</t></p></tc>
          <tc><p><t>교직원 회의</t><lineBreak/><t>운영 협의</t></p></tc>
          <tc><p><t>15:30</t></p></tc>
        </tr></tbl></p></sec>'''
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("Contents/section0.xml", xml)
        event, = parse_schedule_text(extract_file_text(path))
        self.assertEqual((event.title, event.time), ("교직원 회의 운영 협의", "15:30"))

    def test_pdf_independent_columns_keep_each_title_with_its_date(self):
        path = self.root / "columns.pdf"
        write_text_pdf(path, [(40, 735, "2026-11-02 Staff meeting"),
                              (320, 735, "2026-11-06 Club festival"),
                              (40, 695, "2026-11-04 Parent meeting"),
                              (320, 695, "2026-11-09 Safety training")])
        extracted = extract_file_text(path)
        self.assertEqual(extracted.splitlines(), ["2026-11-02 Staff meeting", "2026-11-04 Parent meeting",
                                                  "2026-11-06 Club festival", "2026-11-09 Safety training"])
        events = parse_schedule_text(extracted)
        self.assertEqual(len(events), 4)
        for event, title in zip(events, ("Staff meeting", "Parent meeting", "Club festival", "Safety training")):
            self.assertTrue(event.title.endswith(title), event.title)
            self.assertEqual(sum(other in event.title for other in
                                 ("Staff meeting", "Parent meeting", "Club festival", "Safety training")), 1)

    def test_pdf_start_end_date_columns_do_not_become_independent_lists(self):
        path = self.root / "periods.pdf"
        write_text_pdf(path, [(40, 735, "2026-11-02"), (320, 735, "2026-11-06 Staff camp"),
                              (40, 695, "2026-11-09"), (320, 695, "2026-11-13 Science camp")])
        extracted = extract_file_text(path)
        self.assertEqual(len(extracted.splitlines()), 2)
        self.assertIn("2026-11-02 2026-11-06 Staff camp", extracted)

    def test_pdf_single_two_date_row_is_insufficient_column_evidence(self):
        path = self.root / "single.pdf"
        write_text_pdf(path, [(40, 735, "2026-11-02 First meeting"),
                              (320, 735, "2026-11-06 Second meeting")])
        extracted = extract_file_text(path)
        self.assertEqual(len(extracted.splitlines()), 1)


if __name__ == "__main__":
    unittest.main()
